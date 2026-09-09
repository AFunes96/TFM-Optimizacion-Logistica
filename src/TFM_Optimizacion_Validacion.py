from __future__ import annotations

import copy
import math
import multiprocessing as mp
import os
import platform
import queue as queue_module
import random
import re
import subprocess
import sys
import tempfile
import time
import traceback
from dataclasses import dataclass, field, replace
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Set, Tuple

import numpy as np
import pandas as pd

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt


BASE_DIR = Path(__file__).resolve().parents[1]

RUTA_PRECIOS = BASE_DIR / "data" / "processed" / "enllaç-v2_precios_medios.xlsx"

RUTA_DEMANDA = BASE_DIR / "data" / "generated" / "demanda_sintetica.xlsx"

RUTA_DISTANCIAS = BASE_DIR / "data" / "processed" / "matriz_distancias_120.xlsx"

CARPETA_SALIDA = BASE_DIR / "results" / "Resultados_Validacion_GA"

RUTA_EXCEL_SALIDA = CARPETA_SALIDA / "Resultados_Validacion_GA.xlsx"
CARPETA_GRAFICAS = CARPETA_SALIDA / "Graficas"
CARPETA_MAPAS = CARPETA_SALIDA / "Mapas"
CARPETA_LOGS_MILP = CARPETA_SALIDA / "Logs_MILP"


MODO_EJECUCION = "final"

if MODO_EJECUCION.lower() == "prueba":
    N_SEMANAS_BAJA = 2
    N_SEMANAS_MEDIA = 2
    N_SEMANAS_ALTA = 2
    SEMILLAS_GA = [1, 2, 3]
    N_MILP_POR_NIVEL = 1

    TAMANO_POBLACION = 100
    NUMERO_GENERACIONES = 250
    PACIENCIA = 50

    SEMILLAS_SENSIBILIDAD = [1, 2]
else:
    N_SEMANAS_BAJA = 10
    N_SEMANAS_MEDIA = 10
    N_SEMANAS_ALTA = 10
    SEMILLAS_GA = list(range(1, 11))
    N_MILP_POR_NIVEL = 3

    TAMANO_POBLACION = 200
    NUMERO_GENERACIONES = 500
    PACIENCIA = 100

    SEMILLAS_SENSIBILIDAD = [1, 2, 3]


SEMILLA_SELECCION_SEMANAS = 123


PROBABILIDAD_CRUCE = 0.80
PROBABILIDAD_MUTACION = 0.25
PORCENTAJE_ELITISMO = 0.05
MAX_EXTRAS_INICIALES = 3

NUMERO_MUTACIONES = [1, 2, 3, 4]
PROB_NUMERO_MUTACIONES = [0.50, 0.30, 0.15, 0.05]

TIPOS_MUTACION = [
    "sustitucion",
    "inversion",
    "swap",
    "insercion",
    "anadir",
    "eliminar",
]

PROB_TIPOS_MUTACION = [
    0.30,
    0.25,
    0.15,
    0.10,
    0.10,
    0.10,
]


COSTE_POR_KM_BASE = 1.2747


EJECUTAR_SENSIBILIDAD = True
COSTES_KM_SENSIBILIDAD = [1.00, 1.2747, 1.50]


N_SEMANAS_SENSIBILIDAD_POR_NIVEL = 1


EJECUTAR_MILP = True


MILP_SOLVER = "HiGHS"


TIME_LIMIT_MILP_SEGUNDOS = (
    60 if MODO_EJECUCION.lower() == "prueba" else 300
)


MILP_THREADS = 1
MILP_MOSTRAR_LOG = False
MILP_RANDOM_SEED = 123


MILP_HIGHS_PARALLEL = "off"


MILP_GAP_REL_OBJETIVO = 0.0


MILP_USAR_TIMEOUT_DURO = True


MILP_MARGEN_CIERRE_SOLVER_SEGUNDOS = 3.0

MILP_WATCHDOG_INTERVALO_SEGUNDOS = 10
MILP_MOSTRAR_PROGRESO_WATCHDOG = True


FORMULACION_SUBTOURS_MILP = "MTZ"


LIMITE_TSP_EXACTO_BASELINE = 15


GENERAR_GRAFICAS = True
MOSTRAR_GRAFICAS = False
DPI_GRAFICAS = 180


GENERAR_MAPAS = True
RUTA_DATOS_GEOGRAFICOS = BASE_DIR / "data" / "raw" / "enllaç-v2.xlsx"
MAPAS_MODO_TRAZADO = "lineas_rectas"


HUB_ID = None

PENALIZACION_PRODUCTO_NO_CUBIERTO = 1_000_000_000.0
PENALIZACION_UNIDAD_NO_CUBIERTA = 1_000_000.0

TOLERANCIA_MEJORA = 1e-9


@dataclass
class Individuo:

    orden: List[str]
    k: int

    coste_total: float = math.inf
    coste_compra: float = math.inf
    coste_transporte: float = math.inf
    distancia_km: float = math.inf
    factible: bool = False
    demanda_no_cubierta: float = 0.0

    def activos(self) -> List[str]:
        return self.orden[:self.k]


@dataclass
class DatosMaestros:


    demanda_df: pd.DataFrame
    precios_df: pd.DataFrame
    distancias_df: pd.DataFrame
    distancias_np: np.ndarray
    indice_nodo: Dict[str, int]
    hub: str
    productos_sin_proveedor_df: pd.DataFrame
    resumen_coherencia_df: pd.DataFrame


@dataclass
class DatosProblema:
    escenario: int
    semana: int
    fecha: object
    nivel_demanda: str
    coste_por_km: float

    demanda: Dict[str, float]
    productos: List[str]
    demanda_vector: np.ndarray
    productos_demandados: Set[str]

    productores: List[str]
    productor_a_idx_local: Dict[str, int]
    precio_matriz: np.ndarray

    ofertas_por_producto: Dict[str, Dict[str, float]]
    productos_por_productor: Dict[str, Set[str]]

    hub: str
    distancias_np: np.ndarray
    indice_nodo: Dict[str, int]


    demanda_total: float
    numero_productos_con_demanda: int


    demanda_total_original: float
    demanda_sin_proveedor: float
    pct_demanda_sin_proveedor: float
    productos_con_demanda_originales: int
    productos_sin_proveedor: int


@dataclass
class ResultadoGA:
    mejor: Individuo
    historial: pd.DataFrame
    tiempo_total_s: float
    tiempo_hasta_mejor_s: float
    generacion_mejor: int
    generaciones_ejecutadas: int


def normalizar_id(valor) -> str:
    texto = str(valor).strip()
    if texto.endswith(".0"):
        texto = texto[:-2]
    return texto


def etiqueta_semana(escenario: int, semana: int) -> str:
    return f"E{int(escenario)}-S{int(semana)}"


def asegurar_carpetas():
    CARPETA_SALIDA.mkdir(parents=True, exist_ok=True)
    CARPETA_GRAFICAS.mkdir(parents=True, exist_ok=True)
    CARPETA_MAPAS.mkdir(parents=True, exist_ok=True)
    CARPETA_LOGS_MILP.mkdir(parents=True, exist_ok=True)


def comprobar_archivos_entrada():
    faltantes = []
    for ruta in [RUTA_PRECIOS, RUTA_DEMANDA, RUTA_DISTANCIAS]:
        if not ruta.exists():
            faltantes.append(str(ruta))

    if faltantes:
        raise FileNotFoundError(
            "No se encuentran los siguientes archivos de entrada:\n- "
            + "\n- ".join(faltantes)
        )


def clonar(individuo: Individuo) -> Individuo:
    return copy.deepcopy(individuo)


def sanitizar_nombre_archivo(texto: str) -> str:
    texto = re.sub(r"[^A-Za-z0-9_\-]+", "_", str(texto))
    return texto.strip("_")


def cargar_demanda_completa(ruta: Path) -> pd.DataFrame:
    df = pd.read_excel(ruta, sheet_name="DemandaSintetica")

    necesarias = {
        "Escenario",
        "NumeroSemanaSimulada",
        "IdProducto",
        "Demanda",
    }
    faltan = necesarias - set(df.columns)
    if faltan:
        raise ValueError(
            "Faltan columnas en DemandaSintetica: "
            + ", ".join(sorted(faltan))
        )

    df["IdProducto"] = df["IdProducto"].map(normalizar_id)
    df["Escenario"] = pd.to_numeric(df["Escenario"], errors="raise").astype(int)
    df["NumeroSemanaSimulada"] = pd.to_numeric(
        df["NumeroSemanaSimulada"], errors="raise"
    ).astype(int)
    df["Demanda"] = pd.to_numeric(df["Demanda"], errors="coerce").fillna(0.0)

    if "Fecha" in df.columns:
        df["Fecha"] = pd.to_datetime(df["Fecha"], errors="coerce")

    return df


def cargar_precios(ruta: Path) -> pd.DataFrame:
    df = pd.read_excel(ruta, sheet_name="Productos")

    necesarias = {"IdProducto", "IdProductor", "PrecioPorUnidad"}
    faltan = necesarias - set(df.columns)
    if faltan:
        raise ValueError(
            "Faltan columnas en la hoja Productos: "
            + ", ".join(sorted(faltan))
        )

    df["IdProducto"] = df["IdProducto"].map(normalizar_id)
    df["IdProductor"] = df["IdProductor"].map(normalizar_id)
    df["PrecioPorUnidad"] = pd.to_numeric(
        df["PrecioPorUnidad"], errors="coerce"
    )

    df = df.dropna(subset=["PrecioPorUnidad"])
    df = df[df["PrecioPorUnidad"] >= 0].copy()


    df = (
        df.groupby(["IdProducto", "IdProductor"], as_index=False)
        .agg(PrecioPorUnidad=("PrecioPorUnidad", "mean"))
    )

    return df


def cargar_matriz_distancias(ruta: Path) -> pd.DataFrame:
    df = pd.read_excel(ruta, index_col=0)

    df.index = [normalizar_id(x) for x in df.index]
    df.columns = [normalizar_id(x) for x in df.columns]

    if df.index.duplicated().any():
        raise ValueError("Hay IDs duplicados en filas de la matriz de distancias.")

    if pd.Index(df.columns).duplicated().any():
        raise ValueError("Hay IDs duplicados en columnas de la matriz de distancias.")

    if set(df.index) != set(df.columns):
        raise ValueError(
            "Los IDs de filas y columnas de la matriz de distancias no coinciden."
        )


    df = df.loc[df.index, df.index]
    df = df.apply(pd.to_numeric, errors="coerce")

    if df.isna().any().any():
        raise ValueError("La matriz de distancias contiene valores NaN.")

    return df


def preparar_coherencia_demanda_oferta(
    demanda_df: pd.DataFrame,
    precios_df: pd.DataFrame,
    distancias_df: pd.DataFrame,
    hub: str,
) -> Tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:

    demanda = demanda_df.copy()
    precios = precios_df.copy()

    productores_matriz = set(distancias_df.index) - {hub}
    precios_validos = precios[
        precios["IdProductor"].isin(productores_matriz)
    ].copy()

    productos_en_precios = set(precios["IdProducto"].unique())
    productos_con_proveedor_valido = set(
        precios_validos["IdProducto"].unique()
    )

    demanda["DemandaOriginal"] = pd.to_numeric(
        demanda["Demanda"], errors="coerce"
    ).fillna(0.0)

    demanda["TieneProveedorValido"] = demanda["IdProducto"].isin(
        productos_con_proveedor_valido
    )

    demanda["DemandaSinProveedor"] = np.where(
        demanda["TieneProveedorValido"],
        0.0,
        demanda["DemandaOriginal"],
    )


    demanda["Demanda"] = np.where(
        demanda["TieneProveedorValido"],
        demanda["DemandaOriginal"],
        0.0,
    )

    positivos_sin = demanda[
        (~demanda["TieneProveedorValido"])
        & (demanda["DemandaOriginal"] > 0)
    ].copy()

    if positivos_sin.empty:
        productos_sin = pd.DataFrame(columns=[
            "IdProducto",
            "Motivo",
            "DemandaTotalExcluida",
            "AparicionesPositivas",
            "SemanasConDemanda",
            "EscenariosAfectados",
            "PctDemandaOriginalTotal",
        ])
    else:
        positivos_sin["ClaveSemana"] = (
            positivos_sin["Escenario"].astype(str)
            + "-"
            + positivos_sin["NumeroSemanaSimulada"].astype(str)
        )

        productos_sin = (
            positivos_sin.groupby("IdProducto", as_index=False)
            .agg(
                DemandaTotalExcluida=("DemandaOriginal", "sum"),
                AparicionesPositivas=("DemandaOriginal", "size"),
                SemanasConDemanda=("ClaveSemana", "nunique"),
                EscenariosAfectados=("Escenario", "nunique"),
            )
        )

        def motivo_producto(producto: str) -> str:
            if producto not in productos_en_precios:
                return "Sin proveedor en archivo de precios"
            return "Proveedor(es) en precios pero fuera de la matriz de distancias"

        productos_sin["Motivo"] = productos_sin["IdProducto"].map(
            motivo_producto
        )

        total_original = float(demanda["DemandaOriginal"].sum())
        productos_sin["PctDemandaOriginalTotal"] = np.where(
            total_original > 0,
            productos_sin["DemandaTotalExcluida"] / total_original * 100.0,
            np.nan,
        )

        productos_sin = productos_sin[[
            "IdProducto",
            "Motivo",
            "DemandaTotalExcluida",
            "AparicionesPositivas",
            "SemanasConDemanda",
            "EscenariosAfectados",
            "PctDemandaOriginalTotal",
        ]].sort_values(
            ["DemandaTotalExcluida", "IdProducto"],
            ascending=[False, True],
        ).reset_index(drop=True)

    total_original = float(demanda["DemandaOriginal"].sum())
    total_optimizable = float(demanda["Demanda"].sum())
    total_excluida = float(demanda["DemandaSinProveedor"].sum())

    positivos_originales = demanda[demanda["DemandaOriginal"] > 0]
    positivos_optimizables = demanda[demanda["Demanda"] > 0]

    claves_afectadas = positivos_sin[[
        "Escenario", "NumeroSemanaSimulada"
    ]].drop_duplicates() if not positivos_sin.empty else pd.DataFrame()

    todas_claves = demanda[[
        "Escenario", "NumeroSemanaSimulada"
    ]].drop_duplicates()

    resumen = pd.DataFrame({
        "Indicador": [
            "Productos distintos con demanda positiva original",
            "Productos distintos optimizables con demanda positiva",
            "Productos demandados sin proveedor válido",
            "Demanda total original",
            "Demanda total optimizable",
            "Demanda total excluida por falta de proveedor",
            "Porcentaje de demanda original excluida",
            "Semanas totales en el archivo sintético",
            "Semanas afectadas por demanda sin proveedor",
            "Porcentaje de semanas afectadas",
        ],
        "Valor": [
            int(positivos_originales["IdProducto"].nunique()),
            int(positivos_optimizables["IdProducto"].nunique()),
            int(len(productos_sin)),
            total_original,
            total_optimizable,
            total_excluida,
            (total_excluida / total_original * 100.0) if total_original > 0 else np.nan,
            int(len(todas_claves)),
            int(len(claves_afectadas)),
            (len(claves_afectadas) / len(todas_claves) * 100.0)
                if len(todas_claves) > 0 else np.nan,
        ],
        "Unidad": [
            "productos",
            "productos",
            "productos",
            "unidades",
            "unidades",
            "unidades",
            "%",
            "semanas",
            "semanas",
            "%",
        ],
    })

    return demanda, productos_sin, resumen


def cargar_datos_maestros() -> DatosMaestros:
    comprobar_archivos_entrada()

    print("Cargando demanda sintética...")
    demanda_df = cargar_demanda_completa(RUTA_DEMANDA)

    print("Cargando precios medios...")
    precios_df = cargar_precios(RUTA_PRECIOS)

    print("Cargando matriz de distancias...")
    distancias_df = cargar_matriz_distancias(RUTA_DISTANCIAS)

    if HUB_ID is None:
        hub = normalizar_id(distancias_df.index[-1])
    else:
        hub = normalizar_id(HUB_ID)

    if hub not in distancias_df.index:
        raise ValueError(f"El HUB '{hub}' no aparece en la matriz de distancias.")


    (
        demanda_df,
        productos_sin_proveedor_df,
        resumen_coherencia_df,
    ) = preparar_coherencia_demanda_oferta(
        demanda_df=demanda_df,
        precios_df=precios_df,
        distancias_df=distancias_df,
        hub=hub,
    )

    print("\n" + "=" * 78)
    print("CONTROL DE COHERENCIA DEMANDA-OFERTA")
    print("=" * 78)
    print(resumen_coherencia_df.to_string(index=False))

    if not productos_sin_proveedor_df.empty:
        print(
            "\nAVISO: estos productos se excluyen SOLO del universo de optimización. "
            "La demanda original y la demanda excluida quedan registradas en el Excel final."
        )

    distancias_np = distancias_df.to_numpy(dtype=float)
    indice_nodo = {
        nodo: i for i, nodo in enumerate(distancias_df.index.tolist())
    }

    return DatosMaestros(
        demanda_df=demanda_df,
        precios_df=precios_df,
        distancias_df=distancias_df,
        distancias_np=distancias_np,
        indice_nodo=indice_nodo,
        hub=hub,
        productos_sin_proveedor_df=productos_sin_proveedor_df,
        resumen_coherencia_df=resumen_coherencia_df,
    )


def crear_resumen_semanal(demanda_df: pd.DataFrame) -> pd.DataFrame:

    columnas_grupo = ["Escenario", "NumeroSemanaSimulada"]

    if "Fecha" in demanda_df.columns:
        columnas_grupo.append("Fecha")

    temporal = demanda_df.copy()

    if "DemandaOriginal" not in temporal.columns:
        temporal["DemandaOriginal"] = temporal["Demanda"]
    if "DemandaSinProveedor" not in temporal.columns:
        temporal["DemandaSinProveedor"] = 0.0
    if "TieneProveedorValido" not in temporal.columns:
        temporal["TieneProveedorValido"] = True

    temporal["ConDemandaOriginal"] = (
        temporal["DemandaOriginal"] > 0
    ).astype(int)
    temporal["ConDemandaOptimizable"] = (
        temporal["Demanda"] > 0
    ).astype(int)
    temporal["ProductoSinProveedorConDemanda"] = (
        (temporal["DemandaOriginal"] > 0)
        & (~temporal["TieneProveedorValido"])
    ).astype(int)

    resumen = (
        temporal.groupby(columnas_grupo, as_index=False)
        .agg(
            DemandaTotalOriginal=("DemandaOriginal", "sum"),
            DemandaTotalOptimizable=("Demanda", "sum"),
            DemandaSinProveedor=("DemandaSinProveedor", "sum"),
            ProductosConDemandaOriginales=("ConDemandaOriginal", "sum"),
            ProductosConDemanda=("ConDemandaOptimizable", "sum"),
            ProductosSinProveedor=("ProductoSinProveedorConDemanda", "sum"),
            NumeroProductos=("IdProducto", "nunique"),
        )
    )

    resumen["PctDemandaSinProveedor"] = np.where(
        resumen["DemandaTotalOriginal"] > 0,
        resumen["DemandaSinProveedor"]
        / resumen["DemandaTotalOriginal"] * 100.0,
        0.0,
    )


    resumen["DemandaTotalSemana"] = resumen["DemandaTotalOptimizable"]


    resumen = resumen[resumen["DemandaTotalSemana"] > 0].copy()

    if resumen.empty:
        raise ValueError(
            "No hay ninguna semana con demanda positiva y proveedor válido."
        )

    q33 = float(resumen["DemandaTotalSemana"].quantile(1 / 3))
    q67 = float(resumen["DemandaTotalSemana"].quantile(2 / 3))

    def clasificar(valor):
        if valor <= q33:
            return "Baja"
        if valor <= q67:
            return "Media"
        return "Alta"

    resumen["NivelDemanda"] = resumen["DemandaTotalSemana"].map(clasificar)
    resumen["EtiquetaSemana"] = [
        etiqueta_semana(e, s)
        for e, s in zip(
            resumen["Escenario"],
            resumen["NumeroSemanaSimulada"],
        )
    ]

    resumen.attrs["Q33"] = q33
    resumen.attrs["Q67"] = q67

    return resumen.sort_values(
        ["NivelDemanda", "DemandaTotalSemana", "Escenario", "NumeroSemanaSimulada"]
    ).reset_index(drop=True)


def seleccionar_round_robin_escenarios(
    candidatos: pd.DataFrame,
    n: int,
    rng: random.Random,
) -> pd.DataFrame:

    if n <= 0 or candidatos.empty:
        return candidatos.iloc[0:0].copy()

    grupos = {}
    for escenario, grupo in candidatos.groupby("Escenario"):
        indices = grupo.index.tolist()
        rng.shuffle(indices)
        grupos[int(escenario)] = indices

    escenarios = sorted(grupos)
    rng.shuffle(escenarios)

    seleccionados = []
    usados = set()

    while len(seleccionados) < min(n, len(candidatos)):
        progreso = False

        for escenario in escenarios:
            lista = grupos[escenario]
            while lista and lista[0] in usados:
                lista.pop(0)

            if lista:
                idx = lista.pop(0)
                seleccionados.append(idx)
                usados.add(idx)
                progreso = True

                if len(seleccionados) >= min(n, len(candidatos)):
                    break

        if not progreso:
            break

    return candidatos.loc[seleccionados].copy()


def seleccionar_semanas_experimento(
    resumen: pd.DataFrame,
    semilla: int,
) -> pd.DataFrame:
    rng = random.Random(semilla)

    objetivo = {
        "Baja": N_SEMANAS_BAJA,
        "Media": N_SEMANAS_MEDIA,
        "Alta": N_SEMANAS_ALTA,
    }

    partes = []

    for nivel in ["Baja", "Media", "Alta"]:
        candidatos = resumen[resumen["NivelDemanda"] == nivel].copy()
        n = min(objetivo[nivel], len(candidatos))

        if n < objetivo[nivel]:
            print(
                f"AVISO: solo hay {n} semanas disponibles en nivel {nivel}; "
                f"se solicitaron {objetivo[nivel]}."
            )

        elegido = seleccionar_round_robin_escenarios(candidatos, n, rng)
        partes.append(elegido)

    seleccion = pd.concat(partes, ignore_index=True)

    orden_nivel = pd.Categorical(
        seleccion["NivelDemanda"],
        categories=["Baja", "Media", "Alta"],
        ordered=True,
    )
    seleccion = seleccion.assign(_orden_nivel=orden_nivel)
    seleccion = seleccion.sort_values(
        ["_orden_nivel", "DemandaTotalSemana", "Escenario", "NumeroSemanaSimulada"]
    ).drop(columns="_orden_nivel")

    return seleccion.reset_index(drop=True)


def seleccionar_representativas_por_nivel(
    semanas: pd.DataFrame,
    n_por_nivel: int,
) -> pd.DataFrame:

    partes = []

    for nivel in ["Baja", "Media", "Alta"]:
        grupo = semanas[semanas["NivelDemanda"] == nivel].sort_values(
            "DemandaTotalSemana"
        )

        if grupo.empty:
            continue

        n = min(n_por_nivel, len(grupo))

        if n == 1:
            posiciones = [len(grupo) // 2]
        else:
            posiciones = np.linspace(0, len(grupo) - 1, n + 2)[1:-1]
            posiciones = np.rint(posiciones).astype(int).tolist()


        posiciones_unicas = []
        for pos in posiciones:
            pos = int(np.clip(pos, 0, len(grupo) - 1))
            if pos not in posiciones_unicas:
                posiciones_unicas.append(pos)


        for pos in range(len(grupo)):
            if len(posiciones_unicas) >= n:
                break
            if pos not in posiciones_unicas:
                posiciones_unicas.append(pos)

        partes.append(grupo.iloc[posiciones_unicas[:n]].copy())

    if not partes:
        return semanas.iloc[0:0].copy()

    return pd.concat(partes, ignore_index=True)


def construir_problema_semana(
    maestros: DatosMaestros,
    fila_semana: pd.Series,
    coste_por_km: float = COSTE_POR_KM_BASE,
) -> DatosProblema:

    escenario = int(fila_semana["Escenario"])
    semana = int(fila_semana["NumeroSemanaSimulada"])
    nivel = str(fila_semana["NivelDemanda"])

    filtro = (
        (maestros.demanda_df["Escenario"] == escenario)
        & (maestros.demanda_df["NumeroSemanaSimulada"] == semana)
    )

    df_semana = maestros.demanda_df.loc[filtro].copy()

    if df_semana.empty:
        raise ValueError(f"No se encontró {etiqueta_semana(escenario, semana)}.")

    agregada = (
        df_semana.groupby("IdProducto", as_index=False)["Demanda"]
        .sum()
    )
    agregada = agregada[agregada["Demanda"] > 0].copy()

    if agregada.empty:
        raise ValueError(
            f"{etiqueta_semana(escenario, semana)} tiene demanda total cero."
        )

    productos = agregada["IdProducto"].astype(str).tolist()
    demanda_vector = agregada["Demanda"].astype(float).to_numpy()
    demanda = dict(zip(productos, demanda_vector))
    productos_demandados = set(productos)

    precios = maestros.precios_df[
        maestros.precios_df["IdProducto"].isin(productos_demandados)
    ].copy()


    productores_con_oferta = set(precios["IdProductor"])
    productores = [
        nodo
        for nodo in maestros.distancias_df.index
        if nodo != maestros.hub and nodo in productores_con_oferta
    ]

    precios = precios[precios["IdProductor"].isin(productores)].copy()

    ofertas_por_producto: Dict[str, Dict[str, float]] = {}
    for fila in precios.itertuples(index=False):
        ofertas_por_producto.setdefault(str(fila.IdProducto), {})[
            str(fila.IdProductor)
        ] = float(fila.PrecioPorUnidad)

    sin_oferta = [
        p for p in productos if p not in ofertas_por_producto
        or not ofertas_por_producto[p]
    ]
    if sin_oferta:
        raise ValueError(
            f"{etiqueta_semana(escenario, semana)} no puede satisfacerse. "
            "Productos sin proveedor válido: " + ", ".join(sin_oferta)
        )

    productos_por_productor = {p: set() for p in productores}
    for producto, ofertas in ofertas_por_producto.items():
        for productor in ofertas:
            productos_por_productor[productor].add(producto)

    productor_a_idx_local = {p: i for i, p in enumerate(productores)}

    precio_matriz = np.full(
        (len(productores), len(productos)),
        np.inf,
        dtype=float,
    )

    producto_a_col = {p: j for j, p in enumerate(productos)}
    for producto, ofertas in ofertas_por_producto.items():
        j = producto_a_col[producto]
        for productor, precio in ofertas.items():
            i = productor_a_idx_local[productor]
            precio_matriz[i, j] = float(precio)

    if "Fecha" in df_semana.columns:
        fecha = df_semana["Fecha"].iloc[0]
    else:
        fecha = None

    return DatosProblema(
        escenario=escenario,
        semana=semana,
        fecha=fecha,
        nivel_demanda=nivel,
        coste_por_km=float(coste_por_km),
        demanda=demanda,
        productos=productos,
        demanda_vector=demanda_vector,
        productos_demandados=productos_demandados,
        productores=productores,
        productor_a_idx_local=productor_a_idx_local,
        precio_matriz=precio_matriz,
        ofertas_por_producto=ofertas_por_producto,
        productos_por_productor=productos_por_productor,
        hub=maestros.hub,
        distancias_np=maestros.distancias_np,
        indice_nodo=maestros.indice_nodo,
        demanda_total=float(demanda_vector.sum()),
        numero_productos_con_demanda=len(productos),
        demanda_total_original=float(
            fila_semana.get("DemandaTotalOriginal", demanda_vector.sum())
        ),
        demanda_sin_proveedor=float(
            fila_semana.get("DemandaSinProveedor", 0.0)
        ),
        pct_demanda_sin_proveedor=float(
            fila_semana.get("PctDemandaSinProveedor", 0.0)
        ),
        productos_con_demanda_originales=int(
            fila_semana.get("ProductosConDemandaOriginales", len(productos))
        ),
        productos_sin_proveedor=int(
            fila_semana.get("ProductosSinProveedor", 0)
        ),
    )


def productos_cubiertos(
    productores_activos: Sequence[str],
    datos: DatosProblema,
) -> Set[str]:
    cubiertos: Set[str] = set()
    for productor in productores_activos:
        cubiertos.update(datos.productos_por_productor.get(productor, set()))
    return cubiertos & datos.productos_demandados


def productos_faltantes(individuo: Individuo, datos: DatosProblema) -> Set[str]:
    return datos.productos_demandados - productos_cubiertos(
        individuo.activos(), datos
    )


def reparar_individuo(
    individuo: Individuo,
    datos: DatosProblema,
    rng: random.Random,
) -> Individuo:

    individuo.k = max(1, min(individuo.k, len(individuo.orden)))
    faltantes = productos_faltantes(individuo, datos)

    while faltantes:
        candidatos = [
            productor
            for productor in individuo.orden[individuo.k:]
            if datos.productos_por_productor.get(productor, set()) & faltantes
        ]

        if not candidatos:
            return individuo

        elegido = rng.choice(candidatos)
        pos_antigua = individuo.orden.index(elegido)
        individuo.orden.pop(pos_antigua)

        pos_nueva = rng.randint(0, individuo.k)
        individuo.orden.insert(pos_nueva, elegido)
        individuo.k += 1

        faltantes = productos_faltantes(individuo, datos)

    return individuo


def generar_individuo_inicial(
    datos: DatosProblema,
    rng: random.Random,
) -> Individuo:
    faltantes = set(datos.productos_demandados)
    disponibles = list(datos.productores)
    activos: List[str] = []

    while faltantes:
        candidatos = [
            p for p in disponibles
            if datos.productos_por_productor.get(p, set()) & faltantes
        ]

        if not candidatos:
            raise RuntimeError("No se puede construir un individuo factible.")

        elegido = rng.choice(candidatos)
        disponibles.remove(elegido)
        activos.insert(rng.randint(0, len(activos)), elegido)
        faltantes -= datos.productos_por_productor[elegido]

    max_extras = min(MAX_EXTRAS_INICIALES, len(disponibles))
    n_extras = rng.randint(0, max_extras)

    if n_extras > 0:
        extras = rng.sample(disponibles, n_extras)
        for productor in extras:
            disponibles.remove(productor)
            activos.insert(rng.randint(0, len(activos)), productor)

    rng.shuffle(disponibles)
    return Individuo(orden=activos + disponibles, k=len(activos))


def distancia_ruta_metros(
    productores_activos: Sequence[str],
    datos: DatosProblema,
) -> float:
    nodos = [datos.hub] + list(productores_activos) + [datos.hub]
    indices = [datos.indice_nodo[n] for n in nodos]

    origenes = np.asarray(indices[:-1], dtype=int)
    destinos = np.asarray(indices[1:], dtype=int)
    valores = datos.distancias_np[origenes, destinos]

    if not np.all(np.isfinite(valores)):
        return math.inf

    return float(valores.sum())


def evaluar_individuo(
    individuo: Individuo,
    datos: DatosProblema,
    cache: Optional[dict] = None,
) -> float:

    activos = individuo.activos()
    clave = tuple(activos)

    if cache is not None and clave in cache:
        (
            individuo.coste_total,
            individuo.coste_compra,
            individuo.coste_transporte,
            individuo.distancia_km,
            individuo.factible,
            individuo.demanda_no_cubierta,
        ) = cache[clave]
        return individuo.coste_total

    indices_locales = [datos.productor_a_idx_local[p] for p in activos]

    if not indices_locales:
        individuo.factible = False
        individuo.coste_total = PENALIZACION_PRODUCTO_NO_CUBIERTO
        return individuo.coste_total

    subprecios = datos.precio_matriz[np.asarray(indices_locales, dtype=int), :]
    mejores_precios = np.min(subprecios, axis=0)
    cubierto = np.isfinite(mejores_precios)

    demanda_no_cubierta = float(datos.demanda_vector[~cubierto].sum())
    productos_no_cubiertos = int((~cubierto).sum())

    if cubierto.any():
        coste_compra = float(
            np.dot(datos.demanda_vector[cubierto], mejores_precios[cubierto])
        )
    else:
        coste_compra = 0.0

    distancia_m = distancia_ruta_metros(activos, datos)

    if math.isinf(distancia_m):
        distancia_km = math.inf
        coste_transporte = PENALIZACION_PRODUCTO_NO_CUBIERTO
    else:
        distancia_km = distancia_m / 1000.0
        coste_transporte = distancia_km * datos.coste_por_km

    penalizacion = (
        productos_no_cubiertos * PENALIZACION_PRODUCTO_NO_CUBIERTO
        + demanda_no_cubierta * PENALIZACION_UNIDAD_NO_CUBIERTA
    )

    factible = productos_no_cubiertos == 0 and not math.isinf(distancia_m)
    coste_total = coste_compra + coste_transporte + penalizacion

    individuo.coste_total = float(coste_total)
    individuo.coste_compra = float(coste_compra)
    individuo.coste_transporte = float(coste_transporte)
    individuo.distancia_km = float(distancia_km)
    individuo.factible = bool(factible)
    individuo.demanda_no_cubierta = float(demanda_no_cubierta)

    if cache is not None:
        cache[clave] = (
            individuo.coste_total,
            individuo.coste_compra,
            individuo.coste_transporte,
            individuo.distancia_km,
            individuo.factible,
            individuo.demanda_no_cubierta,
        )

    return individuo.coste_total


def obtener_asignacion_compras(
    productores_activos: Sequence[str],
    datos: DatosProblema,
) -> pd.DataFrame:
    filas = []

    indices_locales = [datos.productor_a_idx_local[p] for p in productores_activos]
    subprecios = datos.precio_matriz[np.asarray(indices_locales, dtype=int), :]

    for j, producto in enumerate(datos.productos):
        columna = subprecios[:, j]
        if not np.isfinite(columna).any():
            continue

        idx_rel = int(np.argmin(columna))
        productor = productores_activos[idx_rel]
        precio = float(columna[idx_rel])
        cantidad = float(datos.demanda_vector[j])

        filas.append({
            "IdProducto": producto,
            "Demanda": cantidad,
            "IdProductor": productor,
            "PrecioUnidad": precio,
            "CosteProducto": cantidad * precio,
        })

    return pd.DataFrame(filas)


def crear_poblacion_inicial(
    datos: DatosProblema,
    rng: random.Random,
    cache: dict,
) -> List[Individuo]:
    poblacion = []

    for _ in range(TAMANO_POBLACION):
        individuo = generar_individuo_inicial(datos, rng)
        reparar_individuo(individuo, datos, rng)
        evaluar_individuo(individuo, datos, cache)

        if not individuo.factible:
            raise RuntimeError("Se generó un individuo inicial no factible.")

        poblacion.append(individuo)

    return poblacion


def preparar_ruleta(poblacion: Sequence[Individuo]):
    ordenados = sorted(poblacion, key=lambda ind: ind.coste_total)
    n = len(ordenados)
    pesos = list(range(n, 0, -1))
    return ordenados, pesos


def seleccionar_padre(
    ordenados: Sequence[Individuo],
    pesos: Sequence[float],
    rng: random.Random,
) -> Individuo:
    return rng.choices(ordenados, weights=pesos, k=1)[0]


def ox_una_direccion(
    orden_a: List[str],
    orden_b: List[str],
    rng: random.Random,
) -> List[str]:
    n = len(orden_a)
    if n < 2:
        return orden_a.copy()

    corte_1, corte_2 = sorted(rng.sample(range(n), 2))

    hijo = [None] * n
    hijo[corte_1:corte_2 + 1] = orden_a[corte_1:corte_2 + 1]
    genes_segmento = set(hijo[corte_1:corte_2 + 1])

    recorrido_b = orden_b[corte_2 + 1:] + orden_b[:corte_2 + 1]
    insertar = [g for g in recorrido_b if g not in genes_segmento]
    posiciones = list(range(corte_2 + 1, n)) + list(range(0, corte_1))

    for pos, gen in zip(posiciones, insertar):
        hijo[pos] = gen

    if any(g is None for g in hijo):
        raise RuntimeError("OX produjo una permutación incompleta.")

    return list(hijo)


def cruzar_ox(
    padre_1: Individuo,
    padre_2: Individuo,
    rng: random.Random,
) -> Tuple[Individuo, Individuo]:
    orden_1 = ox_una_direccion(padre_1.orden, padre_2.orden, rng)
    orden_2 = ox_una_direccion(padre_2.orden, padre_1.orden, rng)

    k_min = min(padre_1.k, padre_2.k)
    k_max = max(padre_1.k, padre_2.k)

    return (
        Individuo(orden_1, rng.randint(k_min, k_max)),
        Individuo(orden_2, rng.randint(k_min, k_max)),
    )


def mutacion_swap(individuo: Individuo, rng: random.Random):
    if individuo.k < 2:
        return
    i, j = rng.sample(range(individuo.k), 2)
    individuo.orden[i], individuo.orden[j] = individuo.orden[j], individuo.orden[i]


def mutacion_inversion(individuo: Individuo, rng: random.Random):
    if individuo.k < 2:
        return
    i, j = sorted(rng.sample(range(individuo.k), 2))
    individuo.orden[i:j + 1] = list(reversed(individuo.orden[i:j + 1]))


def mutacion_insercion(individuo: Individuo, rng: random.Random):
    if individuo.k < 2:
        return
    origen = rng.randrange(individuo.k)
    gen = individuo.orden.pop(origen)
    destino = rng.randrange(individuo.k)
    individuo.orden.insert(destino, gen)


def mutacion_sustitucion(
    individuo: Individuo,
    datos: DatosProblema,
    rng: random.Random,
):
    if individuo.k >= len(individuo.orden):
        return

    pos_activa = rng.randrange(individuo.k)
    saliente = individuo.orden[pos_activa]
    productos_saliente = (
        datos.productos_por_productor.get(saliente, set())
        & datos.productos_demandados
    )

    candidatos = [
        p for p in individuo.orden[individuo.k:]
        if datos.productos_por_productor.get(p, set()) & productos_saliente
    ]

    if not candidatos:
        candidatos = list(individuo.orden[individuo.k:])

    if not candidatos:
        return

    entrante = rng.choice(candidatos)
    pos_inactiva = individuo.orden.index(entrante)

    individuo.orden[pos_activa], individuo.orden[pos_inactiva] = (
        individuo.orden[pos_inactiva],
        individuo.orden[pos_activa],
    )


def mutacion_anadir(
    individuo: Individuo,
    datos: DatosProblema,
    rng: random.Random,
):
    if individuo.k >= len(individuo.orden):
        return

    candidatos = [
        p for p in individuo.orden[individuo.k:]
        if datos.productos_por_productor.get(p, set()) & datos.productos_demandados
    ]

    if not candidatos:
        return

    elegido = rng.choice(candidatos)
    pos_antigua = individuo.orden.index(elegido)
    individuo.orden.pop(pos_antigua)
    individuo.orden.insert(rng.randint(0, individuo.k), elegido)
    individuo.k += 1


def mutacion_eliminar(individuo: Individuo, rng: random.Random):
    if individuo.k <= 1:
        return

    pos = rng.randrange(individuo.k)
    eliminado = individuo.orden.pop(pos)
    individuo.k -= 1
    individuo.orden.insert(
        rng.randint(individuo.k, len(individuo.orden)),
        eliminado,
    )


def mutar(
    individuo: Individuo,
    datos: DatosProblema,
    rng: random.Random,
):
    n_ops = rng.choices(
        NUMERO_MUTACIONES,
        weights=PROB_NUMERO_MUTACIONES,
        k=1,
    )[0]

    for _ in range(n_ops):
        tipo = rng.choices(
            TIPOS_MUTACION,
            weights=PROB_TIPOS_MUTACION,
            k=1,
        )[0]

        if tipo == "sustitucion":
            mutacion_sustitucion(individuo, datos, rng)
        elif tipo == "inversion":
            mutacion_inversion(individuo, rng)
        elif tipo == "swap":
            mutacion_swap(individuo, rng)
        elif tipo == "insercion":
            mutacion_insercion(individuo, rng)
        elif tipo == "anadir":
            mutacion_anadir(individuo, datos, rng)
        elif tipo == "eliminar":
            mutacion_eliminar(individuo, rng)

    reparar_individuo(individuo, datos, rng)


def resumen_generacion(
    poblacion: Sequence[Individuo],
    generacion: int,
    mejor_global: Individuo,
    tiempo_s: float,
) -> dict:
    costes = np.asarray([i.coste_total for i in poblacion], dtype=float)
    mejor_generacion = min(poblacion, key=lambda i: i.coste_total)

    return {
        "Generacion": generacion,
        "MejorCosteGeneracion": mejor_generacion.coste_total,
        "MejorCosteGlobal": mejor_global.coste_total,
        "CosteMedioPoblacion": float(costes.mean()),
        "CosteMedianoPoblacion": float(np.median(costes)),
        "MejorCosteCompra": mejor_global.coste_compra,
        "MejorCosteTransporte": mejor_global.coste_transporte,
        "MejorDistanciaKm": mejor_global.distancia_km,
        "MejorNumeroProductores": mejor_global.k,
        "TiempoAcumuladoSegundos": float(tiempo_s),
    }


def ejecutar_ga(
    datos: DatosProblema,
    semilla_ga: int,
    verbose: bool = False,
) -> ResultadoGA:
    rng = random.Random(int(semilla_ga))
    cache = {}

    t0 = time.perf_counter()

    poblacion = crear_poblacion_inicial(datos, rng, cache)
    mejor_global = clonar(min(poblacion, key=lambda i: i.coste_total))

    generacion_mejor = 0
    tiempo_hasta_mejor = time.perf_counter() - t0
    sin_mejora = 0

    historial = [
        resumen_generacion(
            poblacion,
            0,
            mejor_global,
            time.perf_counter() - t0,
        )
    ]

    if verbose:
        print(
            f"  Gen 0 | coste={mejor_global.coste_total:.2f} | "
            f"km={mejor_global.distancia_km:.2f} | k={mejor_global.k}"
        )

    ultima_generacion = 0

    for generacion in range(1, NUMERO_GENERACIONES + 1):
        ultima_generacion = generacion

        ordenados, pesos = preparar_ruleta(poblacion)

        n_elites = max(
            1,
            int(round(TAMANO_POBLACION * PORCENTAJE_ELITISMO)),
        )

        nueva = [clonar(i) for i in ordenados[:n_elites]]

        while len(nueva) < TAMANO_POBLACION:
            padre_1 = seleccionar_padre(ordenados, pesos, rng)
            padre_2 = seleccionar_padre(ordenados, pesos, rng)

            if rng.random() < PROBABILIDAD_CRUCE:
                hijo_1, hijo_2 = cruzar_ox(padre_1, padre_2, rng)
            else:
                hijo_1 = Individuo(padre_1.orden.copy(), padre_1.k)
                hijo_2 = Individuo(padre_2.orden.copy(), padre_2.k)

            for hijo in [hijo_1, hijo_2]:
                if rng.random() < PROBABILIDAD_MUTACION:
                    mutar(hijo, datos, rng)


                reparar_individuo(hijo, datos, rng)
                evaluar_individuo(hijo, datos, cache)
                nueva.append(hijo)

                if len(nueva) >= TAMANO_POBLACION:
                    break

        poblacion = nueva
        mejor_generacion_actual = min(poblacion, key=lambda i: i.coste_total)

        if (
            mejor_generacion_actual.coste_total
            < mejor_global.coste_total - TOLERANCIA_MEJORA
        ):
            mejor_global = clonar(mejor_generacion_actual)
            generacion_mejor = generacion
            tiempo_hasta_mejor = time.perf_counter() - t0
            sin_mejora = 0
        else:
            sin_mejora += 1

        historial.append(
            resumen_generacion(
                poblacion,
                generacion,
                mejor_global,
                time.perf_counter() - t0,
            )
        )

        if verbose and (generacion == 1 or generacion % 25 == 0):
            print(
                f"  Gen {generacion} | coste={mejor_global.coste_total:.2f} | "
                f"km={mejor_global.distancia_km:.2f} | k={mejor_global.k}"
            )

        if PACIENCIA is not None and sin_mejora >= PACIENCIA:
            break

    tiempo_total = time.perf_counter() - t0


    evaluar_individuo(mejor_global, datos, cache=None)

    return ResultadoGA(
        mejor=mejor_global,
        historial=pd.DataFrame(historial),
        tiempo_total_s=float(tiempo_total),
        tiempo_hasta_mejor_s=float(tiempo_hasta_mejor),
        generacion_mejor=int(generacion_mejor),
        generaciones_ejecutadas=int(ultima_generacion),
    )


def coste_ruta_productores(
    orden_productores: Sequence[str],
    datos: DatosProblema,
) -> float:
    return distancia_ruta_metros(orden_productores, datos) / 1000.0


def tsp_held_karp(
    productores: Sequence[str],
    datos: DatosProblema,
) -> Tuple[List[str], float]:

    productores = list(productores)
    n = len(productores)

    if n == 0:
        return [], 0.0
    if n == 1:
        return productores.copy(), coste_ruta_productores(productores, datos)

    idx_h = datos.indice_nodo[datos.hub]
    idx_p = [datos.indice_nodo[p] for p in productores]


    dp = {}

    for j in range(n):
        mask = 1 << j
        dp[(mask, j)] = (float(datos.distancias_np[idx_h, idx_p[j]]), None)

    for size in range(2, n + 1):

        for mask in range(1, 1 << n):
            if mask.bit_count() != size:
                continue

            for j in range(n):
                if not (mask & (1 << j)):
                    continue

                prev_mask = mask ^ (1 << j)
                mejor_coste = math.inf
                mejor_prev = None

                for i in range(n):
                    if not (prev_mask & (1 << i)):
                        continue

                    clave = (prev_mask, i)
                    if clave not in dp:
                        continue

                    candidato = (
                        dp[clave][0]
                        + float(datos.distancias_np[idx_p[i], idx_p[j]])
                    )

                    if candidato < mejor_coste:
                        mejor_coste = candidato
                        mejor_prev = i

                dp[(mask, j)] = (mejor_coste, mejor_prev)

    full = (1 << n) - 1
    mejor_final = math.inf
    ultimo = None

    for j in range(n):
        coste = dp[(full, j)][0] + float(
            datos.distancias_np[idx_p[j], idx_h]
        )
        if coste < mejor_final:
            mejor_final = coste
            ultimo = j


    orden_indices = []
    mask = full
    j = ultimo

    while j is not None:
        orden_indices.append(j)
        anterior = dp[(mask, j)][1]
        mask ^= 1 << j
        j = anterior

    orden_indices.reverse()
    orden = [productores[j] for j in orden_indices]

    return orden, mejor_final / 1000.0


def ruta_vecino_mas_cercano(
    productores: Sequence[str],
    datos: DatosProblema,
) -> List[str]:
    restantes = set(productores)
    actual = datos.hub
    ruta = []

    while restantes:
        idx_actual = datos.indice_nodo[actual]
        siguiente = min(
            restantes,
            key=lambda p: (
                datos.distancias_np[idx_actual, datos.indice_nodo[p]],
                p,
            ),
        )
        ruta.append(siguiente)
        restantes.remove(siguiente)
        actual = siguiente

    return ruta


def mejorar_2opt_asimetrico(
    ruta: Sequence[str],
    datos: DatosProblema,
    max_iter: int = 100,
) -> Tuple[List[str], float]:

    mejor = list(ruta)
    mejor_dist = coste_ruta_productores(mejor, datos)

    for _ in range(max_iter):
        hubo_mejora = False
        n = len(mejor)

        for i in range(n - 1):
            for j in range(i + 1, n):
                candidata = mejor[:i] + list(reversed(mejor[i:j + 1])) + mejor[j + 1:]
                dist = coste_ruta_productores(candidata, datos)

                if dist < mejor_dist - 1e-12:
                    mejor = candidata
                    mejor_dist = dist
                    hubo_mejora = True
                    break
            if hubo_mejora:
                break

        if not hubo_mejora:
            break

    return mejor, mejor_dist


def ejecutar_baseline(datos: DatosProblema) -> dict:

    t0 = time.perf_counter()

    asignacion = []
    seleccionados = []
    coste_compra = 0.0

    for producto, cantidad in datos.demanda.items():
        ofertas = datos.ofertas_por_producto[producto]
        productor = min(ofertas, key=lambda p: (ofertas[p], p))
        precio = float(ofertas[productor])
        coste = float(cantidad) * precio

        seleccionados.append(productor)
        coste_compra += coste
        asignacion.append({
            "IdProducto": producto,
            "Demanda": float(cantidad),
            "IdProductor": productor,
            "PrecioUnidad": precio,
            "CosteProducto": coste,
        })

    seleccionados = list(dict.fromkeys(seleccionados))

    if len(seleccionados) <= LIMITE_TSP_EXACTO_BASELINE:
        ruta, distancia_km = tsp_held_karp(seleccionados, datos)
        metodo_ruta = "Held-Karp exacto"
    else:
        inicial = ruta_vecino_mas_cercano(seleccionados, datos)
        ruta, distancia_km = mejorar_2opt_asimetrico(inicial, datos)
        metodo_ruta = "Vecino más cercano + 2-opt"

    coste_transporte = distancia_km * datos.coste_por_km
    coste_total = coste_compra + coste_transporte
    tiempo_s = time.perf_counter() - t0

    return {
        "Escenario": datos.escenario,
        "NumeroSemanaSimulada": datos.semana,
        "NivelDemanda": datos.nivel_demanda,
        "EtiquetaSemana": etiqueta_semana(datos.escenario, datos.semana),
        "CosteTotalBaseline": coste_total,
        "CosteCompraBaseline": coste_compra,
        "CosteTransporteBaseline": coste_transporte,
        "DistanciaKmBaseline": distancia_km,
        "NumeroProductoresBaseline": len(seleccionados),
        "MetodoRutaBaseline": metodo_ruta,
        "TiempoBaselineSegundos": tiempo_s,
        "RutaBaseline": " -> ".join([datos.hub] + ruta + [datos.hub]),
        "AsignacionBaseline": asignacion,
    }


def _parsear_log_highs(ruta_log: Path) -> dict:
    resultado = {
        "MILP_OptimoDemostrado": False,
        "MILP_LimiteTiempo": False,
        "MILP_TerminoPorGap": False,
        "MILP_LowerBound": np.nan,
        "MILP_PrimalBound": np.nan,
        "MILP_GapReportadoPct": np.nan,
        "MILP_NodosExplorados": np.nan,
        "MILP_ResultText": "",
    }

    if not ruta_log.exists():
        return resultado

    try:
        texto = ruta_log.read_text(encoding="utf-8", errors="ignore")
    except Exception:
        return resultado

    texto_lower = texto.lower()


    m_status = re.search(r"Status\s+([^\n\r]+)", texto, flags=re.IGNORECASE)
    status_text = m_status.group(1).strip() if m_status else ""
    status_lower = status_text.lower()

    if "time limit reached" in status_lower:
        resultado["MILP_LimiteTiempo"] = True
        resultado["MILP_ResultText"] = "Time limit reached"
    elif "optimal" in status_lower:
        resultado["MILP_ResultText"] = status_text


    def extraer_float(patron):
        m = re.search(patron, texto, flags=re.IGNORECASE)
        if not m:
            return np.nan
        valor = m.group(1).strip()
        if valor.lower() in {"inf", "+inf", "-inf", "infinity", "+infinity", "-infinity"}:
            return np.nan
        try:
            return float(valor)
        except ValueError:
            return np.nan

    resultado["MILP_PrimalBound"] = extraer_float(
        r"Primal bound\s+([-+0-9.eE]+|[-+]?inf(?:inity)?)"
    )
    resultado["MILP_LowerBound"] = extraer_float(
        r"Dual bound\s+([-+0-9.eE]+|[-+]?inf(?:inity)?)"
    )

    m_gap = re.search(
        r"Gap\s+([-+0-9.eE]+|inf)\s*%?",
        texto,
        flags=re.IGNORECASE,
    )
    if m_gap:
        try:
            if m_gap.group(1).lower() != "inf":
                resultado["MILP_GapReportadoPct"] = abs(float(m_gap.group(1)))
        except ValueError:
            pass

    m_nodes = re.search(r"Nodes\s+(\d+)", texto, flags=re.IGNORECASE)
    if m_nodes:
        try:
            resultado["MILP_NodosExplorados"] = int(m_nodes.group(1))
        except ValueError:
            pass


    if "optimal" in status_lower:
        if MILP_GAP_REL_OBJETIVO <= 0:
            resultado["MILP_OptimoDemostrado"] = True
            resultado["MILP_TerminoPorGap"] = False
        else:
            gap_pct = resultado["MILP_GapReportadoPct"]
            if np.isfinite(gap_pct) and gap_pct > 1e-7:
                resultado["MILP_TerminoPorGap"] = (
                    gap_pct <= MILP_GAP_REL_OBJETIVO * 100.0 + 1e-6
                )
            else:
                resultado["MILP_OptimoDemostrado"] = True


    if "time limit reached" in texto_lower:
        resultado["MILP_LimiteTiempo"] = True

    return resultado


def _extraer_info_highs(problema, status_pulp: str) -> dict:
    info = {
        "MILP_StatusSolver": status_pulp,
        "MILP_OptimoDemostrado": False,
        "MILP_LimiteTiempo": False,
        "MILP_TerminoPorGap": False,
        "MILP_LowerBound": np.nan,
        "MILP_PrimalBound": np.nan,
        "MILP_GapReportadoPct": np.nan,
        "MILP_NodosExplorados": np.nan,
        "MILP_ResultText": "",
    }

    modelo = getattr(problema, "solverModel", None)
    if modelo is None:
        return info

    try:
        estado = modelo.getModelStatus()
        estado_texto = str(modelo.modelStatusToString(estado))
        info["MILP_StatusSolver"] = estado_texto
        info["MILP_ResultText"] = estado_texto
    except Exception:
        estado_texto = status_pulp

    try:
        highs_info = modelo.getInfo()
    except Exception:
        highs_info = None

    if highs_info is not None:
        try:
            v = float(getattr(highs_info, "mip_dual_bound"))
            if np.isfinite(v):
                info["MILP_LowerBound"] = v
        except Exception:
            pass

        try:
            v = float(getattr(highs_info, "objective_function_value"))
            if np.isfinite(v):
                info["MILP_PrimalBound"] = v
        except Exception:
            pass

        try:
            v = float(getattr(highs_info, "mip_gap"))
            if np.isfinite(v):
                info["MILP_GapReportadoPct"] = max(0.0, v * 100.0)
        except Exception:
            pass

        try:
            info["MILP_NodosExplorados"] = int(
                getattr(highs_info, "mip_node_count")
            )
        except Exception:
            pass

    estado_lower = str(info["MILP_StatusSolver"]).lower()
    gap_pct = info["MILP_GapReportadoPct"]

    if "time limit" in estado_lower:
        info["MILP_LimiteTiempo"] = True

    if "optimal" in estado_lower:
        if MILP_GAP_REL_OBJETIVO <= 0:
            info["MILP_OptimoDemostrado"] = True
            info["MILP_TerminoPorGap"] = False
        elif np.isfinite(gap_pct) and gap_pct > 1e-7:
            info["MILP_TerminoPorGap"] = (
                gap_pct <= MILP_GAP_REL_OBJETIVO * 100.0 + 1e-6
            )
        else:
            info["MILP_OptimoDemostrado"] = True

    return info


def _resolver_milp_directo(datos: DatosProblema) -> Tuple[dict, pd.DataFrame, pd.DataFrame]:

    try:
        import pulp
    except ImportError as exc:
        raise ImportError(
            "Para ejecutar el MILP instala PuLP y HiGHS: "
            "python -m pip install -U \"pulp[highs]\""
        ) from exc


    try:
        import highspy
    except ImportError as exc:
        raise ImportError(
            "HiGHS no está instalado en este entorno. Ejecuta en la consola "
            "del mismo Python que usa Spyder:\n"
            "python -m pip install -U highspy"
        ) from exc

    if not hasattr(pulp, "HiGHS"):
        raise RuntimeError(
            "Tu versión de PuLP no expone pulp.HiGHS. Actualiza PuLP con:\n"
            "python -m pip install -U \"pulp[highs]\""
        )

    etiqueta = etiqueta_semana(datos.escenario, datos.semana)
    log_path = CARPETA_LOGS_MILP / f"HIGHS_{etiqueta}.log"

    try:
        if log_path.exists():
            log_path.unlink()
    except OSError:
        pass

    t_total_ini = time.perf_counter()

    problema = pulp.LpProblem(f"TFM_{etiqueta}", pulp.LpMinimize)

    productores = list(datos.productores)
    productos = list(datos.productos)
    hub = datos.hub
    nodos = [hub] + productores
    n_prod = len(productores)

    if n_prod == 0:
        raise ValueError(
            f"{etiqueta}: no hay productores candidatos para el MILP."
        )


    y = {
        i: pulp.LpVariable(f"y_{i}", cat="Binary")
        for i in productores
    }

    x = {}
    for p in productos:
        for i in datos.ofertas_por_producto[p]:
            x[(i, p)] = pulp.LpVariable(f"x_{i}_{p}", cat="Binary")

    arcos = [(i, j) for i in nodos for j in nodos if i != j]
    z = {
        (i, j): pulp.LpVariable(f"z_{i}_{j}", cat="Binary")
        for i, j in arcos
    }

    u = {
        i: pulp.LpVariable(
            f"u_{i}", lowBound=0, upBound=n_prod, cat="Continuous"
        )
        for i in productores
    }


    coste_compra = pulp.lpSum(
        datos.demanda[p]
        * datos.ofertas_por_producto[p][i]
        * x[(i, p)]
        for (i, p) in x
    )

    coste_transporte = pulp.lpSum(
        (
            datos.distancias_np[
                datos.indice_nodo[i],
                datos.indice_nodo[j],
            ]
            / 1000.0
        )
        * datos.coste_por_km
        * z[(i, j)]
        for i, j in arcos
    )

    problema += coste_compra + coste_transporte


    for p in productos:
        problema += (
            pulp.lpSum(x[(i, p)] for i in datos.ofertas_por_producto[p]) == 1,
            f"Asignacion_{p}",
        )
        for i in datos.ofertas_por_producto[p]:
            problema += (
                x[(i, p)] <= y[i],
                f"AsignacionImplicaVisita_{i}_{p}",
            )


    for i in productores:
        problema += (
            pulp.lpSum(z[(i, j)] for j in nodos if j != i) == y[i],
            f"Salida_{i}",
        )
        problema += (
            pulp.lpSum(z[(j, i)] for j in nodos if j != i) == y[i],
            f"Entrada_{i}",
        )

    problema += (
        pulp.lpSum(z[(hub, j)] for j in productores) == 1,
        "Salida_HUB",
    )
    problema += (
        pulp.lpSum(z[(i, hub)] for i in productores) == 1,
        "Entrada_HUB",
    )


    for i in productores:
        problema += (u[i] >= y[i], f"MTZ_LB_{i}")
        problema += (u[i] <= n_prod * y[i], f"MTZ_UB_{i}")

    for i in productores:
        for j in productores:
            if i == j:
                continue
            problema += (
                u[i] - u[j] + n_prod * z[(i, j)]
                <= n_prod - y[j],
                f"MTZ_{i}_{j}",
            )

    tiempo_construccion = time.perf_counter() - t_total_ini
    tiempo_restante = TIME_LIMIT_MILP_SEGUNDOS - tiempo_construccion

    n_variables = len(problema.variables())
    n_restricciones = len(problema.constraints)

    print(
        f"  Productores candidatos: {n_prod} | "
        f"Productos demandados: {len(productos)}"
    )
    print(
        f"  Variables: {n_variables} "
        f"(binarias aprox.: {len(y) + len(x) + len(z)}; MTZ: {len(u)}) | "
        f"Restricciones: {n_restricciones}"
    )
    print(f"  Construcción MILP: {tiempo_construccion:.2f} s")

    if tiempo_restante <= 0:
        tiempo_total = time.perf_counter() - t_total_ini
        print(
            f"  Presupuesto agotado antes de HiGHS "
            f"({TIME_LIMIT_MILP_SEGUNDOS:.0f} s por instancia)."
        )

        fila_resultado = {
            "Escenario": datos.escenario,
            "NumeroSemanaSimulada": datos.semana,
            "NivelDemanda": datos.nivel_demanda,
            "EtiquetaSemana": etiqueta,
            "DemandaTotalSemana": datos.demanda_total,
            "DemandaTotalOriginal": datos.demanda_total_original,
            "DemandaSinProveedor": datos.demanda_sin_proveedor,
            "PctDemandaSinProveedor": datos.pct_demanda_sin_proveedor,
            "ProductosConDemandaOriginales": datos.productos_con_demanda_originales,
            "ProductosConDemanda": datos.numero_productos_con_demanda,
            "ProductosSinProveedor": datos.productos_sin_proveedor,
            "ProductoresCandidatosMILP": n_prod,
            "MILP_Solver": MILP_SOLVER,
            "MILP_FormulacionSubtours": FORMULACION_SUBTOURS_MILP,
            "MILP_NumeroVariables": n_variables,
            "MILP_NumeroRestricciones": n_restricciones,
            "MILP_StatusPuLP": "TIME_LIMIT_CONSTRUCCION",
            "MILP_StatusSolver": "TIME_LIMIT_CONSTRUCCION",
            "MILP_OptimoDemostrado": False,
            "MILP_LimiteTiempo": True,
            "MILP_TerminoPorGap": False,
            "MILP_TimeoutDuro": False,
            "MILP_GapObjetivoPct": MILP_GAP_REL_OBJETIVO * 100.0,
            "MILP_CosteTotal": np.nan,
            "MILP_CosteReconstruido": np.nan,
            "MILP_CosteCompra": np.nan,
            "MILP_CosteTransporte": np.nan,
            "MILP_DistanciaKm": np.nan,
            "MILP_NumeroProductores": np.nan,
            "MILP_PrimalBound": np.nan,
            "MILP_LowerBound": np.nan,
            "MILP_GapReportadoPct": np.nan,
            "MILP_GapCalculadoPct": np.nan,
            "MILP_NodosExplorados": np.nan,
            "MILP_TiempoConstruccionSegundos": tiempo_construccion,
            "MILP_TiempoSolverSegundos": 0.0,
            "MILP_TiempoTotalSegundos": tiempo_total,
            "MILP_PresupuestoSolverSegundos": 0.0,
            "MILP_RutaValida": False,
            "MILP_AsignacionCompleta": False,
            "MILP_Ruta": "",
            "MILP_LogPath": str(log_path),
            "MILP_Error": "Presupuesto temporal agotado durante la construcción.",
        }
        return fila_resultado, pd.DataFrame(), pd.DataFrame()

    margen_cierre = (
        MILP_MARGEN_CIERRE_SOLVER_SEGUNDOS
        if MILP_USAR_TIMEOUT_DURO
        else 0.0
    )
    limite_solver = max(1.0, float(tiempo_restante - margen_cierre))

    print(
        f"  HiGHS iniciado | límite restante: {limite_solver:.2f} s | "
        f"objetivo=óptimo exacto | threads={MILP_THREADS}"
    )

    solver = pulp.HiGHS(
        msg=MILP_MOSTRAR_LOG,
        timeLimit=limite_solver,
        gapRel=MILP_GAP_REL_OBJETIVO,
        threads=MILP_THREADS,
        logPath=str(log_path),
        random_seed=MILP_RANDOM_SEED,
        parallel=MILP_HIGHS_PARALLEL,
    )

    try:
        disponible = solver.available()
    except Exception:
        disponible = True

    if disponible is False:
        raise RuntimeError(
            "HiGHS no está disponible. Instálalo con:\n"
            "python -m pip install -U \"pulp[highs]\""
        )

    t_solver_ini = time.perf_counter()
    problema.solve(solver)
    tiempo_solver = time.perf_counter() - t_solver_ini
    tiempo_total = time.perf_counter() - t_total_ini

    status_pulp = pulp.LpStatus.get(problema.status, str(problema.status))
    info_highs = _extraer_info_highs(problema, status_pulp)


    info_log = _parsear_log_highs(log_path)
    for clave in (
        "MILP_LowerBound",
        "MILP_PrimalBound",
        "MILP_GapReportadoPct",
        "MILP_NodosExplorados",
    ):
        if not np.isfinite(info_highs.get(clave, np.nan)):
            valor = info_log.get(clave, np.nan)
            if np.isfinite(valor):
                info_highs[clave] = valor

    if not info_highs["MILP_LimiteTiempo"]:
        info_highs["MILP_LimiteTiempo"] = info_log["MILP_LimiteTiempo"]
    if not info_highs["MILP_TerminoPorGap"]:
        info_highs["MILP_TerminoPorGap"] = info_log["MILP_TerminoPorGap"]
    if not info_highs["MILP_OptimoDemostrado"]:
        info_highs["MILP_OptimoDemostrado"] = info_log["MILP_OptimoDemostrado"]

    print(
        f"  HiGHS finalizado | status={info_highs['MILP_StatusSolver']} | "
        f"solver={tiempo_solver:.2f} s | total={tiempo_total:.2f} s | "
        f"gap={info_highs['MILP_GapReportadoPct'] if np.isfinite(info_highs['MILP_GapReportadoPct']) else np.nan:.4g}%"
    )

    valor_objetivo = pulp.value(problema.objective)
    if valor_objetivo is None or not np.isfinite(valor_objetivo):
        valor_objetivo = info_highs.get("MILP_PrimalBound", np.nan)
    if valor_objetivo is None or not np.isfinite(valor_objetivo):
        valor_objetivo = np.nan


    activos = [
        i for i in productores
        if y[i].value() is not None and y[i].value() > 0.5
    ]

    arcos_usados = [
        (i, j)
        for i, j in arcos
        if z[(i, j)].value() is not None and z[(i, j)].value() > 0.5
    ]

    siguiente = {i: j for i, j in arcos_usados}
    ruta_productores = []
    actual = hub
    visitados = set()

    while actual in siguiente:
        prox = siguiente[actual]
        if prox == hub:
            break
        if prox in visitados:
            break
        ruta_productores.append(prox)
        visitados.add(prox)
        actual = prox

    ruta_valida = (
        len(activos) > 0
        and len(ruta_productores) == len(activos)
        and set(ruta_productores) == set(activos)
    )

    filas_asignacion = []
    coste_compra_val = 0.0

    for p in productos:
        elegido = None
        for i in datos.ofertas_por_producto[p]:
            valor = x[(i, p)].value()
            if valor is not None and valor > 0.5:
                elegido = i
                break

        if elegido is not None:
            cantidad = float(datos.demanda[p])
            precio = float(datos.ofertas_por_producto[p][elegido])
            coste = cantidad * precio
            coste_compra_val += coste
            filas_asignacion.append({
                "Escenario": datos.escenario,
                "NumeroSemanaSimulada": datos.semana,
                "IdProducto": p,
                "Demanda": cantidad,
                "IdProductor": elegido,
                "PrecioUnidad": precio,
                "CosteProducto": coste,
            })

    asignacion_completa = len(filas_asignacion) == len(productos)

    if ruta_valida:
        distancia_km = coste_ruta_productores(ruta_productores, datos)
        coste_transporte_val = distancia_km * datos.coste_por_km
    else:
        distancia_km = np.nan
        coste_transporte_val = np.nan


    if ruta_valida and asignacion_completa:
        coste_reconstruido = coste_compra_val + coste_transporte_val
        if not np.isfinite(valor_objetivo):
            valor_objetivo = coste_reconstruido
    else:
        coste_reconstruido = np.nan

    lower_bound = info_highs.get("MILP_LowerBound", np.nan)
    if (
        np.isfinite(valor_objetivo)
        and np.isfinite(lower_bound)
        and abs(valor_objetivo) > 1e-12
    ):
        gap_calc = max(
            0.0,
            float((valor_objetivo - lower_bound) / abs(valor_objetivo) * 100.0),
        )
    else:
        gap_calc = np.nan


    if info_highs["MILP_OptimoDemostrado"] and np.isfinite(valor_objetivo):
        lower_bound = valor_objetivo
        info_highs["MILP_GapReportadoPct"] = 0.0
        gap_calc = 0.0


    if (
        not info_highs["MILP_OptimoDemostrado"]
        and not info_highs["MILP_TerminoPorGap"]
        and tiempo_solver >= max(0.0, limite_solver - 1.0)
    ):
        info_highs["MILP_LimiteTiempo"] = True

    fila_resultado = {
        "Escenario": datos.escenario,
        "NumeroSemanaSimulada": datos.semana,
        "NivelDemanda": datos.nivel_demanda,
        "EtiquetaSemana": etiqueta,
        "DemandaTotalSemana": datos.demanda_total,
        "DemandaTotalOriginal": datos.demanda_total_original,
        "DemandaSinProveedor": datos.demanda_sin_proveedor,
        "PctDemandaSinProveedor": datos.pct_demanda_sin_proveedor,
        "ProductosConDemandaOriginales": datos.productos_con_demanda_originales,
        "ProductosConDemanda": datos.numero_productos_con_demanda,
        "ProductosSinProveedor": datos.productos_sin_proveedor,
        "ProductoresCandidatosMILP": n_prod,
        "MILP_Solver": MILP_SOLVER,
        "MILP_FormulacionSubtours": FORMULACION_SUBTOURS_MILP,
        "MILP_NumeroVariables": n_variables,
        "MILP_NumeroRestricciones": n_restricciones,
        "MILP_StatusPuLP": status_pulp,
        "MILP_StatusSolver": info_highs["MILP_StatusSolver"],
        "MILP_OptimoDemostrado": info_highs["MILP_OptimoDemostrado"],
        "MILP_LimiteTiempo": info_highs["MILP_LimiteTiempo"],
        "MILP_TerminoPorGap": info_highs["MILP_TerminoPorGap"],
        "MILP_TimeoutDuro": False,
        "MILP_GapObjetivoPct": MILP_GAP_REL_OBJETIVO * 100.0,
        "MILP_CosteTotal": valor_objetivo,
        "MILP_CosteReconstruido": coste_reconstruido,
        "MILP_CosteCompra": coste_compra_val if asignacion_completa else np.nan,
        "MILP_CosteTransporte": coste_transporte_val,
        "MILP_DistanciaKm": distancia_km,
        "MILP_NumeroProductores": len(activos) if ruta_valida else np.nan,
        "MILP_PrimalBound": info_highs.get("MILP_PrimalBound", np.nan),
        "MILP_LowerBound": lower_bound,
        "MILP_GapReportadoPct": info_highs.get("MILP_GapReportadoPct", np.nan),
        "MILP_GapCalculadoPct": gap_calc,
        "MILP_NodosExplorados": info_highs.get("MILP_NodosExplorados", np.nan),
        "MILP_TiempoConstruccionSegundos": tiempo_construccion,
        "MILP_TiempoSolverSegundos": tiempo_solver,
        "MILP_TiempoTotalSegundos": tiempo_total,
        "MILP_PresupuestoSolverSegundos": limite_solver,
        "MILP_RutaValida": ruta_valida,
        "MILP_AsignacionCompleta": asignacion_completa,
        "MILP_Ruta": (
            " -> ".join([hub] + ruta_productores + [hub])
            if ruta_valida else ""
        ),
        "MILP_LogPath": str(log_path),
        "MILP_Error": "",
    }

    df_arcos = pd.DataFrame([
        {
            "Escenario": datos.escenario,
            "NumeroSemanaSimulada": datos.semana,
            "Origen": i,
            "Destino": j,
        }
        for i, j in arcos_usados
    ])

    return fila_resultado, pd.DataFrame(filas_asignacion), df_arcos


def _worker_resolver_milp(datos: DatosProblema, cola):
    try:
        resultado, asignacion, arcos = _resolver_milp_directo(datos)
        cola.put(("OK", resultado, asignacion, arcos))
    except BaseException as exc:
        cola.put((
            "ERROR",
            type(exc).__name__,
            str(exc),
            traceback.format_exc(),
        ))


def _ultima_linea_log(ruta_log: Path) -> str:
    try:
        if not ruta_log.exists():
            return ""
        texto = ruta_log.read_text(encoding="utf-8", errors="ignore")
        lineas = [linea.strip() for linea in texto.splitlines() if linea.strip()]
        return lineas[-1] if lineas else ""
    except Exception:
        return ""


def _matar_arbol_proceso(proceso) -> None:
    if proceso is None or proceso.pid is None:
        return

    if not proceso.is_alive():
        proceso.join(timeout=1)
        return

    if os.name == "nt":
        try:
            subprocess.run(
                ["taskkill", "/PID", str(proceso.pid), "/T", "/F"],
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
                check=False,
                timeout=10,
            )
        except Exception:
            try:
                proceso.terminate()
            except Exception:
                pass
    else:
        try:
            proceso.terminate()
        except Exception:
            pass

    proceso.join(timeout=5)

    if proceso.is_alive():
        try:
            proceso.kill()
        except Exception:
            pass
        proceso.join(timeout=2)


def _resultado_timeout_duro(datos: DatosProblema, tiempo_total: float, log_path: Path) -> dict:
    info_log = _parsear_log_highs(log_path)

    return {
        "Escenario": datos.escenario,
        "NumeroSemanaSimulada": datos.semana,
        "NivelDemanda": datos.nivel_demanda,
        "EtiquetaSemana": etiqueta_semana(datos.escenario, datos.semana),
        "DemandaTotalSemana": datos.demanda_total,
        "DemandaTotalOriginal": datos.demanda_total_original,
        "DemandaSinProveedor": datos.demanda_sin_proveedor,
        "PctDemandaSinProveedor": datos.pct_demanda_sin_proveedor,
        "ProductosConDemandaOriginales": datos.productos_con_demanda_originales,
        "ProductosConDemanda": datos.numero_productos_con_demanda,
        "ProductosSinProveedor": datos.productos_sin_proveedor,
        "ProductoresCandidatosMILP": len(datos.productores),
        "MILP_Solver": MILP_SOLVER,
        "MILP_FormulacionSubtours": FORMULACION_SUBTOURS_MILP,
        "MILP_NumeroVariables": np.nan,
        "MILP_NumeroRestricciones": np.nan,
        "MILP_StatusPuLP": "TIMEOUT_DURO",
        "MILP_StatusSolver": "TIMEOUT_DURO",
        "MILP_OptimoDemostrado": False,
        "MILP_LimiteTiempo": True,
        "MILP_TerminoPorGap": False,
        "MILP_TimeoutDuro": True,
        "MILP_GapObjetivoPct": MILP_GAP_REL_OBJETIVO * 100.0,
        "MILP_CosteTotal": np.nan,
        "MILP_CosteReconstruido": np.nan,
        "MILP_CosteCompra": np.nan,
        "MILP_CosteTransporte": np.nan,
        "MILP_DistanciaKm": np.nan,
        "MILP_NumeroProductores": np.nan,
        "MILP_PrimalBound": info_log.get("MILP_PrimalBound", np.nan),
        "MILP_LowerBound": info_log.get("MILP_LowerBound", np.nan),
        "MILP_GapReportadoPct": info_log.get("MILP_GapReportadoPct", np.nan),
        "MILP_GapCalculadoPct": np.nan,
        "MILP_NodosExplorados": info_log.get("MILP_NodosExplorados", np.nan),
        "MILP_TiempoConstruccionSegundos": np.nan,
        "MILP_TiempoSolverSegundos": np.nan,
        "MILP_TiempoTotalSegundos": float(tiempo_total),
        "MILP_PresupuestoSolverSegundos": TIME_LIMIT_MILP_SEGUNDOS,
        "MILP_RutaValida": False,
        "MILP_AsignacionCompleta": False,
        "MILP_Ruta": "",
        "MILP_LogPath": str(log_path),
        "MILP_Error": (
            "Timeout duro externo: HiGHS/PuLP no devolvió el control dentro de "
            f"{TIME_LIMIT_MILP_SEGUNDOS:.0f} s. Se finalizó el árbol de procesos."
        ),
    }


def resolver_milp(datos: DatosProblema) -> Tuple[dict, pd.DataFrame, pd.DataFrame]:
    if not MILP_USAR_TIMEOUT_DURO:
        return _resolver_milp_directo(datos)

    etiqueta = etiqueta_semana(datos.escenario, datos.semana)
    log_path = CARPETA_LOGS_MILP / f"HIGHS_{etiqueta}.log"


    contexto = mp.get_context("spawn")
    cola = contexto.Queue()
    proceso = contexto.Process(
        target=_worker_resolver_milp,
        args=(datos, cola),
        name=f"MILP_{etiqueta}",
    )

    inicio = time.perf_counter()
    proceso.start()

    siguiente_aviso = float(MILP_WATCHDOG_INTERVALO_SEGUNDOS)
    limite = float(TIME_LIMIT_MILP_SEGUNDOS)

    while proceso.is_alive():
        transcurrido = time.perf_counter() - inicio
        restante = limite - transcurrido

        if restante <= 0:
            break

        proceso.join(timeout=min(1.0, restante))

        transcurrido = time.perf_counter() - inicio
        if (
            MILP_MOSTRAR_PROGRESO_WATCHDOG
            and transcurrido >= siguiente_aviso
            and proceso.is_alive()
        ):
            ultima = _ultima_linea_log(log_path)
            mensaje_log = f" | HiGHS: {ultima}" if ultima else ""
            print(
                f"  Watchdog MILP: {transcurrido:.0f}/{limite:.0f} s"
                f"{mensaje_log}"
            )
            siguiente_aviso += float(MILP_WATCHDOG_INTERVALO_SEGUNDOS)

    tiempo_total_externo = time.perf_counter() - inicio

    if proceso.is_alive():
        print(
            f"  TIMEOUT DURO: {etiqueta} superó {limite:.0f} s. "
            "Finalizando HiGHS y continuando..."
        )
        _matar_arbol_proceso(proceso)
        try:
            cola.close()
            cola.join_thread()
        except Exception:
            pass
        return (
            _resultado_timeout_duro(datos, tiempo_total_externo, log_path),
            pd.DataFrame(),
            pd.DataFrame(),
        )

    proceso.join(timeout=1)

    try:
        mensaje = cola.get(timeout=3)
    except queue_module.Empty:
        try:
            cola.close()
            cola.join_thread()
        except Exception:
            pass
        raise RuntimeError(
            f"{etiqueta}: el proceso MILP terminó sin devolver resultado. "
            f"Código de salida: {proceso.exitcode}. Revisa {log_path}."
        )
    finally:
        try:
            cola.close()
            cola.join_thread()
        except Exception:
            pass

    if mensaje[0] == "ERROR":
        _, tipo_error, texto_error, traza = mensaje
        raise RuntimeError(
            f"{etiqueta}: error en worker MILP [{tipo_error}]: {texto_error}\n{traza}"
        )

    _, resultado, asignacion, arcos = mensaje
    resultado["MILP_TimeoutDuro"] = False
    resultado["MILP_GapObjetivoPct"] = MILP_GAP_REL_OBJETIVO * 100.0
    resultado["MILP_TiempoExternoWatchdogSegundos"] = tiempo_total_externo

    return resultado, asignacion, arcos


def ejecutar_bateria_ga(
    maestros: DatosMaestros,
    semanas: pd.DataFrame,
) -> Tuple[pd.DataFrame, pd.DataFrame, dict, pd.DataFrame, pd.DataFrame]:
    filas_ejecuciones = []
    historiales = []
    mejores_por_semana = {}
    rutas_mejores = []
    compras_mejores = []

    total = len(semanas) * len(SEMILLAS_GA)
    contador = 0

    for _, fila_semana in semanas.iterrows():
        datos = construir_problema_semana(
            maestros,
            fila_semana,
            coste_por_km=COSTE_POR_KM_BASE,
        )

        mejor_semana: Optional[ResultadoGA] = None
        mejor_semilla = None

        for semilla in SEMILLAS_GA:
            contador += 1
            print(
                f"GA {contador}/{total} | "
                f"{etiqueta_semana(datos.escenario, datos.semana)} | "
                f"{datos.nivel_demanda} | semilla={semilla}"
            )

            resultado = ejecutar_ga(datos, semilla_ga=semilla, verbose=False)
            mejor = resultado.mejor

            filas_ejecuciones.append({
                "Escenario": datos.escenario,
                "NumeroSemanaSimulada": datos.semana,
                "Fecha": datos.fecha,
                "NivelDemanda": datos.nivel_demanda,
                "EtiquetaSemana": etiqueta_semana(datos.escenario, datos.semana),
                "SemillaGA": semilla,
                "DemandaTotalSemana": datos.demanda_total,
                "DemandaTotalOriginal": datos.demanda_total_original,
                "DemandaSinProveedor": datos.demanda_sin_proveedor,
                "PctDemandaSinProveedor": datos.pct_demanda_sin_proveedor,
                "ProductosConDemandaOriginales": datos.productos_con_demanda_originales,
                "ProductosConDemanda": datos.numero_productos_con_demanda,
                "ProductosSinProveedor": datos.productos_sin_proveedor,
                "ProductoresCandidatos": len(datos.productores),
                "Factible": mejor.factible,
                "CosteTotalGA": mejor.coste_total,
                "CosteCompraGA": mejor.coste_compra,
                "CosteTransporteGA": mejor.coste_transporte,
                "DistanciaKmGA": mejor.distancia_km,
                "NumeroProductoresGA": mejor.k,
                "GeneracionMejor": resultado.generacion_mejor,
                "GeneracionesEjecutadas": resultado.generaciones_ejecutadas,
                "TiempoTotalGASegundos": resultado.tiempo_total_s,
                "TiempoHastaMejorSegundos": resultado.tiempo_hasta_mejor_s,
                "RutaGA": " -> ".join([datos.hub] + mejor.activos() + [datos.hub]),
            })

            hist = resultado.historial.copy()
            hist.insert(0, "SemillaGA", semilla)
            hist.insert(0, "EtiquetaSemana", etiqueta_semana(datos.escenario, datos.semana))
            hist.insert(0, "NivelDemanda", datos.nivel_demanda)
            hist.insert(0, "NumeroSemanaSimulada", datos.semana)
            hist.insert(0, "Escenario", datos.escenario)
            historiales.append(hist)

            if (
                mejor_semana is None
                or mejor.coste_total < mejor_semana.mejor.coste_total
            ):
                mejor_semana = ResultadoGA(
                    mejor=clonar(mejor),
                    historial=resultado.historial.copy(),
                    tiempo_total_s=resultado.tiempo_total_s,
                    tiempo_hasta_mejor_s=resultado.tiempo_hasta_mejor_s,
                    generacion_mejor=resultado.generacion_mejor,
                    generaciones_ejecutadas=resultado.generaciones_ejecutadas,
                )
                mejor_semilla = semilla

        assert mejor_semana is not None

        clave = (datos.escenario, datos.semana)
        mejores_por_semana[clave] = {
            "resultado": mejor_semana,
            "semilla": mejor_semilla,
            "datos": datos,
        }


        ruta = [datos.hub] + mejor_semana.mejor.activos() + [datos.hub]
        acumulada = 0.0
        for tramo, (origen, destino) in enumerate(zip(ruta[:-1], ruta[1:]), start=1):
            d_km = float(
                datos.distancias_np[
                    datos.indice_nodo[origen],
                    datos.indice_nodo[destino],
                ]
            ) / 1000.0
            acumulada += d_km
            rutas_mejores.append({
                "Escenario": datos.escenario,
                "NumeroSemanaSimulada": datos.semana,
                "NivelDemanda": datos.nivel_demanda,
                "SemillaMejorGA": mejor_semilla,
                "Tramo": tramo,
                "Origen": origen,
                "Destino": destino,
                "DistanciaKm": d_km,
                "DistanciaAcumuladaKm": acumulada,
            })

        asignacion = obtener_asignacion_compras(
            mejor_semana.mejor.activos(), datos
        )
        if not asignacion.empty:
            asignacion.insert(0, "SemillaMejorGA", mejor_semilla)
            asignacion.insert(0, "NivelDemanda", datos.nivel_demanda)
            asignacion.insert(0, "NumeroSemanaSimulada", datos.semana)
            asignacion.insert(0, "Escenario", datos.escenario)
            compras_mejores.append(asignacion)

    df_ejecuciones = pd.DataFrame(filas_ejecuciones)
    df_historial = pd.concat(historiales, ignore_index=True) if historiales else pd.DataFrame()
    df_rutas = pd.DataFrame(rutas_mejores)
    df_compras = pd.concat(compras_mejores, ignore_index=True) if compras_mejores else pd.DataFrame()

    return df_ejecuciones, df_historial, mejores_por_semana, df_rutas, df_compras


def crear_resumen_ga_por_semana(ejecuciones: pd.DataFrame) -> pd.DataFrame:
    df = ejecuciones.copy()

    mejor_por_semana = df.groupby(
        ["Escenario", "NumeroSemanaSimulada"]
    )["CosteTotalGA"].transform("min")

    df["GapRespectoMejorSemanaPct"] = np.where(
        mejor_por_semana > 0,
        (df["CosteTotalGA"] - mejor_por_semana) / mejor_por_semana * 100.0,
        np.nan,
    )

    df["Dentro1PctMejor"] = df["GapRespectoMejorSemanaPct"] <= 1.0 + 1e-12

    resumen = (
        df.groupby(
            [
                "Escenario",
                "NumeroSemanaSimulada",
                "NivelDemanda",
                "EtiquetaSemana",
            ],
            as_index=False,
        )
        .agg(
            DemandaTotalSemana=("DemandaTotalSemana", "first"),
            DemandaTotalOriginal=("DemandaTotalOriginal", "first"),
            DemandaSinProveedor=("DemandaSinProveedor", "first"),
            PctDemandaSinProveedor=("PctDemandaSinProveedor", "first"),
            ProductosConDemandaOriginales=("ProductosConDemandaOriginales", "first"),
            ProductosConDemanda=("ProductosConDemanda", "first"),
            ProductosSinProveedor=("ProductosSinProveedor", "first"),
            EjecucionesGA=("SemillaGA", "count"),
            FactibilidadPct=("Factible", "mean"),
            CosteGAMedio=("CosteTotalGA", "mean"),
            CosteGAMediano=("CosteTotalGA", "median"),
            MejorCosteGA=("CosteTotalGA", "min"),
            PeorCosteGA=("CosteTotalGA", "max"),
            StdCosteGA=("CosteTotalGA", "std"),
            DistanciaMediaKm=("DistanciaKmGA", "mean"),
            MejorDistanciaKm=("DistanciaKmGA", "min"),
            ProductoresMedia=("NumeroProductoresGA", "mean"),
            TiempoGAMedioSegundos=("TiempoTotalGASegundos", "mean"),
            TiempoGAMedianoSegundos=("TiempoTotalGASegundos", "median"),
            TiempoHastaMejorMedioSegundos=("TiempoHastaMejorSegundos", "mean"),
            GeneracionMejorMedia=("GeneracionMejor", "mean"),
            PorcentajeDentro1PctMejor=("Dentro1PctMejor", "mean"),
            GapMedioRespectoMejorPct=("GapRespectoMejorSemanaPct", "mean"),
        )
    )

    resumen["FactibilidadPct"] *= 100.0
    resumen["PorcentajeDentro1PctMejor"] *= 100.0
    resumen["CV_CostePct"] = np.where(
        resumen["CosteGAMedio"].abs() > 1e-12,
        resumen["StdCosteGA"] / resumen["CosteGAMedio"].abs() * 100.0,
        np.nan,
    )


    ejecuciones.drop(
        columns=[c for c in ["GapRespectoMejorSemanaPct", "Dentro1PctMejor"] if c in ejecuciones.columns],
        inplace=True,
        errors="ignore",
    )
    ejecuciones["GapRespectoMejorSemanaPct"] = df["GapRespectoMejorSemanaPct"].values
    ejecuciones["Dentro1PctMejor"] = df["Dentro1PctMejor"].values

    return resumen


def crear_resumen_por_nivel(resumen_semana: pd.DataFrame) -> pd.DataFrame:
    return (
        resumen_semana.groupby("NivelDemanda", as_index=False)
        .agg(
            NumeroSemanas=("EtiquetaSemana", "count"),
            DemandaMedia=("DemandaTotalSemana", "mean"),
            DemandaOriginalMedia=("DemandaTotalOriginal", "mean"),
            DemandaSinProveedorMedia=("DemandaSinProveedor", "mean"),
            PctDemandaSinProveedorMedio=("PctDemandaSinProveedor", "mean"),
            ProductosDemandadosMedia=("ProductosConDemanda", "mean"),
            ProductosSinProveedorMedia=("ProductosSinProveedor", "mean"),
            CosteGAMedio=("CosteGAMedio", "mean"),
            MejorCosteGAMedio=("MejorCosteGA", "mean"),
            CV_CosteMedioPct=("CV_CostePct", "mean"),
            Dentro1PctMedioPct=("PorcentajeDentro1PctMejor", "mean"),
            FactibilidadMediaPct=("FactibilidadPct", "mean"),
            TiempoGAMedioSegundos=("TiempoGAMedioSegundos", "mean"),
            DistanciaMediaKm=("DistanciaMediaKm", "mean"),
            ProductoresMedia=("ProductoresMedia", "mean"),
        )
    )


def ejecutar_baselines(
    maestros: DatosMaestros,
    semanas: pd.DataFrame,
) -> pd.DataFrame:
    filas = []

    print("\nCalculando baseline de todas las semanas...")

    for _, fila in semanas.iterrows():
        datos = construir_problema_semana(
            maestros,
            fila,
            coste_por_km=COSTE_POR_KM_BASE,
        )
        filas.append(ejecutar_baseline(datos))

    df = pd.DataFrame(filas)

    if "AsignacionBaseline" in df.columns:
        df = df.drop(columns="AsignacionBaseline")
    return df


def comparar_ga_baseline(
    resumen_ga: pd.DataFrame,
    baseline: pd.DataFrame,
) -> pd.DataFrame:
    df = resumen_ga.merge(
        baseline,
        on=[
            "Escenario",
            "NumeroSemanaSimulada",
            "NivelDemanda",
            "EtiquetaSemana",
        ],
        how="left",
    )

    df["MejoraMejorGA_vs_BaselinePct"] = np.where(
        df["CosteTotalBaseline"] > 0,
        (df["CosteTotalBaseline"] - df["MejorCosteGA"])
        / df["CosteTotalBaseline"]
        * 100.0,
        np.nan,
    )

    df["MejoraGAMedio_vs_BaselinePct"] = np.where(
        df["CosteTotalBaseline"] > 0,
        (df["CosteTotalBaseline"] - df["CosteGAMedio"])
        / df["CosteTotalBaseline"]
        * 100.0,
        np.nan,
    )

    return df


def ejecutar_milps(
    maestros: DatosMaestros,
    semanas_milp: pd.DataFrame,
) -> Tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    if not EJECUTAR_MILP or semanas_milp.empty:
        return pd.DataFrame(), pd.DataFrame(), pd.DataFrame()

    resultados = []
    asignaciones = []
    arcos = []

    print("\nEjecutando MILP de referencia...")

    for idx, (_, fila) in enumerate(semanas_milp.iterrows(), start=1):
        datos = construir_problema_semana(
            maestros,
            fila,
            coste_por_km=COSTE_POR_KM_BASE,
        )

        print(
            f"MILP {idx}/{len(semanas_milp)} | "
            f"{etiqueta_semana(datos.escenario, datos.semana)} | "
            f"{datos.nivel_demanda}"
        )

        try:
            res, asig, arc = resolver_milp(datos)
            resultados.append(res)
            if not asig.empty:
                asignaciones.append(asig)
            if not arc.empty:
                arcos.append(arc)
        except Exception as exc:
            resultados.append({
                "Escenario": datos.escenario,
                "NumeroSemanaSimulada": datos.semana,
                "NivelDemanda": datos.nivel_demanda,
                "EtiquetaSemana": etiqueta_semana(datos.escenario, datos.semana),
                "DemandaTotalSemana": datos.demanda_total,
                "DemandaTotalOriginal": datos.demanda_total_original,
                "DemandaSinProveedor": datos.demanda_sin_proveedor,
                "PctDemandaSinProveedor": datos.pct_demanda_sin_proveedor,
                "ProductosConDemandaOriginales": datos.productos_con_demanda_originales,
                "ProductosConDemanda": datos.numero_productos_con_demanda,
                "ProductosSinProveedor": datos.productos_sin_proveedor,
                "ProductoresCandidatosMILP": len(datos.productores),
                "MILP_Solver": MILP_SOLVER,
        "MILP_FormulacionSubtours": FORMULACION_SUBTOURS_MILP,
                "MILP_NumeroVariables": np.nan,
                "MILP_NumeroRestricciones": np.nan,
                "MILP_StatusPuLP": "ERROR",
                "MILP_StatusSolver": "ERROR",
                "MILP_OptimoDemostrado": False,
                "MILP_LimiteTiempo": False,
                "MILP_TerminoPorGap": False,
                "MILP_TimeoutDuro": False,
                "MILP_GapObjetivoPct": MILP_GAP_REL_OBJETIVO * 100.0,
                "MILP_CosteTotal": np.nan,
                "MILP_CosteReconstruido": np.nan,
                "MILP_CosteCompra": np.nan,
                "MILP_CosteTransporte": np.nan,
                "MILP_DistanciaKm": np.nan,
                "MILP_NumeroProductores": np.nan,
                "MILP_PrimalBound": np.nan,
                "MILP_LowerBound": np.nan,
                "MILP_GapReportadoPct": np.nan,
                "MILP_GapCalculadoPct": np.nan,
                "MILP_NodosExplorados": np.nan,
                "MILP_TiempoConstruccionSegundos": np.nan,
                "MILP_TiempoSolverSegundos": np.nan,
                "MILP_TiempoTotalSegundos": np.nan,
                "MILP_PresupuestoSolverSegundos": np.nan,
                "MILP_RutaValida": False,
                "MILP_AsignacionCompleta": False,
                "MILP_Ruta": "",
                "MILP_LogPath": str(CARPETA_LOGS_MILP / f"HIGHS_{etiqueta_semana(datos.escenario, datos.semana)}.log"),
                "MILP_Error": str(exc),
            })
            print("  ERROR MILP:", exc)

    return (
        pd.DataFrame(resultados),
        pd.concat(asignaciones, ignore_index=True) if asignaciones else pd.DataFrame(),
        pd.concat(arcos, ignore_index=True) if arcos else pd.DataFrame(),
    )


def calcular_tiempo_hasta_umbral(
    historial: pd.DataFrame,
    coste_objetivo: float,
    tolerancia_pct: float = 1.0,
) -> float:
    if not np.isfinite(coste_objetivo) or coste_objetivo <= 0:
        return np.nan

    umbral = coste_objetivo * (1.0 + tolerancia_pct / 100.0)
    cumple = historial[historial["MejorCosteGlobal"] <= umbral]

    if cumple.empty:
        return np.nan

    return float(cumple.iloc[0]["TiempoAcumuladoSegundos"])


def comparar_ga_milp(
    resumen_ga: pd.DataFrame,
    ejecuciones_ga: pd.DataFrame,
    historial_ga: pd.DataFrame,
    milp: pd.DataFrame,
) -> Tuple[pd.DataFrame, pd.DataFrame]:
    if milp.empty:
        return pd.DataFrame(), pd.DataFrame()

    df = resumen_ga.merge(
        milp,
        on=[
            "Escenario",
            "NumeroSemanaSimulada",
            "NivelDemanda",
            "EtiquetaSemana",
        ],
        how="inner",
    )

    df["GapMejorGA_vs_MILPPct"] = np.where(
        df["MILP_CosteTotal"] > 0,
        (df["MejorCosteGA"] - df["MILP_CosteTotal"])
        / df["MILP_CosteTotal"]
        * 100.0,
        np.nan,
    )

    df["GapGAMedio_vs_MILPPct"] = np.where(
        df["MILP_CosteTotal"] > 0,
        (df["CosteGAMedio"] - df["MILP_CosteTotal"])
        / df["MILP_CosteTotal"]
        * 100.0,
        np.nan,
    )

    df["SpeedupMILP_sobre_GA"] = np.where(
        df["TiempoGAMedioSegundos"] > 0,
        df["MILP_TiempoTotalSegundos"] / df["TiempoGAMedioSegundos"],
        np.nan,
    )


    filas_tiempo = []

    for _, fila_m in milp.iterrows():
        coste_ref = fila_m.get("MILP_CosteTotal", np.nan)
        e = int(fila_m["Escenario"])
        s = int(fila_m["NumeroSemanaSimulada"])

        sub_ejec = ejecuciones_ga[
            (ejecuciones_ga["Escenario"] == e)
            & (ejecuciones_ga["NumeroSemanaSimulada"] == s)
        ]

        for _, ejec in sub_ejec.iterrows():
            semilla = int(ejec["SemillaGA"])
            hist = historial_ga[
                (historial_ga["Escenario"] == e)
                & (historial_ga["NumeroSemanaSimulada"] == s)
                & (historial_ga["SemillaGA"] == semilla)
            ]

            t1 = calcular_tiempo_hasta_umbral(hist, coste_ref, tolerancia_pct=1.0)

            filas_tiempo.append({
                "Escenario": e,
                "NumeroSemanaSimulada": s,
                "NivelDemanda": fila_m["NivelDemanda"],
                "EtiquetaSemana": fila_m["EtiquetaSemana"],
                "SemillaGA": semilla,
                "MILP_CosteReferencia": coste_ref,
                "MILP_OptimoDemostrado": fila_m.get("MILP_OptimoDemostrado", False),
                "TiempoGA_Hasta1PctMILP_Segundos": t1,
                "Alcanza1PctMILP": bool(np.isfinite(t1)),
            })

    tiempos = pd.DataFrame(filas_tiempo)

    return df, tiempos


def ejecutar_sensibilidad(
    maestros: DatosMaestros,
    semanas_representativas: pd.DataFrame,
) -> pd.DataFrame:
    if not EJECUTAR_SENSIBILIDAD or semanas_representativas.empty:
        return pd.DataFrame()

    filas = []
    total = (
        len(semanas_representativas)
        * len(COSTES_KM_SENSIBILIDAD)
        * len(SEMILLAS_SENSIBILIDAD)
    )
    contador = 0

    print("\nEjecutando sensibilidad al coste por km...")

    for _, fila_semana in semanas_representativas.iterrows():
        for coste_km in COSTES_KM_SENSIBILIDAD:
            for semilla in SEMILLAS_SENSIBILIDAD:
                contador += 1
                datos = construir_problema_semana(
                    maestros,
                    fila_semana,
                    coste_por_km=float(coste_km),
                )

                print(
                    f"Sensibilidad {contador}/{total} | "
                    f"{etiqueta_semana(datos.escenario, datos.semana)} | "
                    f"{coste_km:.4f} €/km | seed={semilla}"
                )

                res = ejecutar_ga(datos, semilla_ga=semilla, verbose=False)

                filas.append({
                    "Escenario": datos.escenario,
                    "NumeroSemanaSimulada": datos.semana,
                    "NivelDemanda": datos.nivel_demanda,
                    "EtiquetaSemana": etiqueta_semana(datos.escenario, datos.semana),
                    "CostePorKm": coste_km,
                    "SemillaGA": semilla,
                    "CosteTotalGA": res.mejor.coste_total,
                    "CosteCompraGA": res.mejor.coste_compra,
                    "CosteTransporteGA": res.mejor.coste_transporte,
                    "DistanciaKmGA": res.mejor.distancia_km,
                    "NumeroProductoresGA": res.mejor.k,
                    "TiempoGASegundos": res.tiempo_total_s,
                })

    return pd.DataFrame(filas)


def crear_resumen_final(
    ejecuciones: pd.DataFrame,
    resumen_semanas: pd.DataFrame,
    comparacion_baseline: pd.DataFrame,
    comparacion_milp: pd.DataFrame,
    tiempos_1pct: pd.DataFrame,
    resumen_coherencia: pd.DataFrame,
) -> pd.DataFrame:
    filas = []

    def add(indicador, valor, unidad=""):
        filas.append({"Indicador": indicador, "Valor": valor, "Unidad": unidad})


    if resumen_coherencia is not None and not resumen_coherencia.empty:
        for fila in resumen_coherencia.itertuples(index=False):
            add(str(fila.Indicador), fila.Valor, str(fila.Unidad))

    add("Semanas analizadas con GA", len(resumen_semanas), "semanas")
    add("Ejecuciones totales del GA", len(ejecuciones), "ejecuciones")
    add("Factibilidad global GA", ejecuciones["Factible"].mean() * 100.0, "%")
    add("Coste GA medio global", ejecuciones["CosteTotalGA"].mean(), "€")
    add("Tiempo GA medio por ejecución", ejecuciones["TiempoTotalGASegundos"].mean(), "s")
    add("Generación media del mejor", ejecuciones["GeneracionMejor"].mean(), "generaciones")
    add("CV medio semanal del coste", resumen_semanas["CV_CostePct"].mean(), "%")
    add(
        "Ejecuciones dentro del 1 % del mejor semanal",
        resumen_semanas["PorcentajeDentro1PctMejor"].mean(),
        "%",
    )

    if not comparacion_baseline.empty:
        add(
            "Mejora media del mejor GA frente al baseline",
            comparacion_baseline["MejoraMejorGA_vs_BaselinePct"].mean(),
            "%",
        )
        add(
            "Mejora media del GA medio frente al baseline",
            comparacion_baseline["MejoraGAMedio_vs_BaselinePct"].mean(),
            "%",
        )

    if not comparacion_milp.empty:
        n_milp = len(comparacion_milp)
        n_opt = int(comparacion_milp["MILP_OptimoDemostrado"].fillna(False).sum())
        add("Instancias MILP ejecutadas", n_milp, "instancias")
        add("MILP con óptimo demostrado", n_opt, "instancias")

        solo_opt = comparacion_milp[
            comparacion_milp["MILP_OptimoDemostrado"] == True
        ]
        if not solo_opt.empty:
            add(
                "Gap medio del mejor GA respecto al óptimo MILP",
                solo_opt["GapMejorGA_vs_MILPPct"].mean(),
                "%",
            )
            add(
                "Gap medio del GA medio respecto al óptimo MILP",
                solo_opt["GapGAMedio_vs_MILPPct"].mean(),
                "%",
            )

        add(
            "Speedup medio MILP/GA",
            comparacion_milp["SpeedupMILP_sobre_GA"].mean(),
            "veces",
        )

    if not tiempos_1pct.empty:
        add(
            "Ejecuciones GA que alcanzan <=1 % del coste MILP",
            tiempos_1pct["Alcanza1PctMILP"].mean() * 100.0,
            "%",
        )
        add(
            "Tiempo medio GA hasta <=1 % del coste MILP",
            tiempos_1pct["TiempoGA_Hasta1PctMILP_Segundos"].mean(),
            "s",
        )

    return pd.DataFrame(filas)


def guardar_figura(fig, nombre: str):
    ruta = CARPETA_GRAFICAS / nombre
    fig.tight_layout()
    fig.savefig(ruta, dpi=DPI_GRAFICAS, bbox_inches="tight")

    if MOSTRAR_GRAFICAS:
        plt.show()

    plt.close(fig)


def grafica_convergencia_representativas(
    semanas_rep: pd.DataFrame,
    mejores_por_semana: dict,
):
    for _, fila in semanas_rep.iterrows():
        clave = (int(fila["Escenario"]), int(fila["NumeroSemanaSimulada"]))
        if clave not in mejores_por_semana:
            continue

        info = mejores_por_semana[clave]
        hist = info["resultado"].historial
        nivel = str(fila["NivelDemanda"])
        etiqueta = str(fila["EtiquetaSemana"])

        fig, ax = plt.subplots(figsize=(9, 5))
        ax.plot(hist["Generacion"], hist["MejorCosteGlobal"], label="Mejor coste global")
        ax.plot(hist["Generacion"], hist["CosteMedioPoblacion"], label="Coste medio población")
        ax.set_title(f"Convergencia GA - {nivel} - {etiqueta}")
        ax.set_xlabel("Generación")
        ax.set_ylabel("Coste (€)")
        ax.grid(True, alpha=0.25)
        ax.legend()

        guardar_figura(
            fig,
            f"01_convergencia_{sanitizar_nombre_archivo(nivel)}_{etiqueta}.png",
        )


def grafica_estabilidad(ejecuciones: pd.DataFrame):
    niveles = [n for n in ["Baja", "Media", "Alta"] if n in set(ejecuciones["NivelDemanda"])]
    datos_box = [
        ejecuciones.loc[
            ejecuciones["NivelDemanda"] == nivel,
            "GapRespectoMejorSemanaPct",
        ].dropna().to_numpy()
        for nivel in niveles
    ]

    if not datos_box:
        return

    fig, ax = plt.subplots(figsize=(8, 5))
    ax.boxplot(datos_box, tick_labels=niveles)
    ax.set_title("Estabilidad estocástica del GA")
    ax.set_xlabel("Nivel de demanda")
    ax.set_ylabel("Gap respecto al mejor de la semana (%)")
    ax.grid(True, axis="y", alpha=0.25)
    guardar_figura(fig, "02_estabilidad_estocastica.png")


def grafica_mejora_baseline(comparacion: pd.DataFrame):
    if comparacion.empty:
        return

    df = comparacion.sort_values(
        ["NivelDemanda", "DemandaTotalSemana"]
    ).reset_index(drop=True)

    fig, ax = plt.subplots(figsize=(10, max(6, len(df) * 0.25)))
    ax.barh(df["EtiquetaSemana"], df["MejoraMejorGA_vs_BaselinePct"])
    ax.axvline(0, linewidth=1)
    ax.set_title("Mejora del mejor GA frente al baseline")
    ax.set_xlabel("Mejora de coste (%)")
    ax.set_ylabel("Semana")
    ax.grid(True, axis="x", alpha=0.25)
    guardar_figura(fig, "03_mejora_GA_vs_baseline.png")


def grafica_costes_ga_baseline(comparacion: pd.DataFrame):
    if comparacion.empty:
        return

    df = comparacion.sort_values("DemandaTotalSemana").reset_index(drop=True)
    x = np.arange(len(df))
    ancho = 0.40

    fig, ax = plt.subplots(figsize=(12, 6))
    ax.bar(x - ancho / 2, df["MejorCosteGA"], width=ancho, label="Mejor GA")
    ax.bar(x + ancho / 2, df["CosteTotalBaseline"], width=ancho, label="Baseline")
    ax.set_title("Coste total: GA frente al baseline")
    ax.set_xlabel("Semanas ordenadas por demanda total")
    ax.set_ylabel("Coste (€)")
    ax.set_xticks(x)
    ax.set_xticklabels(df["EtiquetaSemana"], rotation=90)
    ax.legend()
    ax.grid(True, axis="y", alpha=0.25)
    guardar_figura(fig, "04_coste_GA_vs_baseline.png")


def grafica_gap_milp(comparacion: pd.DataFrame):
    if comparacion.empty:
        return

    df = comparacion[comparacion["MILP_OptimoDemostrado"] == True].copy()
    if df.empty:
        return

    x = np.arange(len(df))
    ancho = 0.40

    fig, ax = plt.subplots(figsize=(10, 5))
    ax.bar(x - ancho / 2, df["GapMejorGA_vs_MILPPct"], width=ancho, label="Mejor GA")
    ax.bar(x + ancho / 2, df["GapGAMedio_vs_MILPPct"], width=ancho, label="GA medio")
    ax.axhline(0, linewidth=1)
    ax.set_title("Gap del GA respecto al óptimo MILP")
    ax.set_xlabel("Instancia")
    ax.set_ylabel("Gap (%)")
    ax.set_xticks(x)
    ax.set_xticklabels(df["EtiquetaSemana"], rotation=45, ha="right")
    ax.legend()
    ax.grid(True, axis="y", alpha=0.25)
    guardar_figura(fig, "05_gap_GA_vs_MILP.png")


def grafica_tiempos_ga_milp(comparacion: pd.DataFrame):
    if comparacion.empty:
        return

    df = comparacion.dropna(
        subset=["MILP_TiempoTotalSegundos", "TiempoGAMedioSegundos"]
    ).copy()
    if df.empty:
        return

    x = np.arange(len(df))
    ancho = 0.40

    fig, ax = plt.subplots(figsize=(10, 5))
    ax.bar(x - ancho / 2, df["TiempoGAMedioSegundos"], width=ancho, label="GA medio")
    ax.bar(x + ancho / 2, df["MILP_TiempoTotalSegundos"], width=ancho, label="MILP")
    ax.set_title("Tiempo computacional: GA frente a MILP")
    ax.set_xlabel("Instancia")
    ax.set_ylabel("Tiempo (s)")
    ax.set_xticks(x)
    ax.set_xticklabels(df["EtiquetaSemana"], rotation=45, ha="right")
    ax.legend()
    ax.grid(True, axis="y", alpha=0.25)
    guardar_figura(fig, "06_tiempo_GA_vs_MILP.png")


def grafica_productos_distancia(resumen: pd.DataFrame):
    if resumen.empty:
        return

    fig, ax = plt.subplots(figsize=(8, 5))
    for nivel in ["Baja", "Media", "Alta"]:
        sub = resumen[resumen["NivelDemanda"] == nivel]
        if not sub.empty:
            ax.scatter(
                sub["ProductosConDemanda"],
                sub["MejorDistanciaKm"],
                label=nivel,
            )

    ax.set_title("Complejidad de demanda y distancia de la solución")
    ax.set_xlabel("Productos con demanda")
    ax.set_ylabel("Distancia del mejor GA (km)")
    ax.legend()
    ax.grid(True, alpha=0.25)
    guardar_figura(fig, "07_productos_vs_distancia.png")


def grafica_componentes_coste(ejecuciones: pd.DataFrame):
    if ejecuciones.empty:
        return

    mejor_idx = ejecuciones.groupby(
        ["Escenario", "NumeroSemanaSimulada"]
    )["CosteTotalGA"].idxmin()
    mejores = ejecuciones.loc[mejor_idx].copy()

    resumen = (
        mejores.groupby("NivelDemanda", as_index=False)
        .agg(
            CosteCompra=("CosteCompraGA", "mean"),
            CosteTransporte=("CosteTransporteGA", "mean"),
        )
    )

    niveles = [n for n in ["Baja", "Media", "Alta"] if n in set(resumen["NivelDemanda"])]
    resumen = resumen.set_index("NivelDemanda").loc[niveles].reset_index()

    fig, ax = plt.subplots(figsize=(8, 5))
    ax.bar(resumen["NivelDemanda"], resumen["CosteCompra"], label="Compra")
    ax.bar(
        resumen["NivelDemanda"],
        resumen["CosteTransporte"],
        bottom=resumen["CosteCompra"],
        label="Transporte",
    )
    ax.set_title("Composición del coste del mejor GA por nivel de demanda")
    ax.set_xlabel("Nivel de demanda")
    ax.set_ylabel("Coste medio (€)")
    ax.legend()
    ax.grid(True, axis="y", alpha=0.25)
    guardar_figura(fig, "08_componentes_coste_por_nivel.png")


def grafica_sensibilidad(sensibilidad: pd.DataFrame):
    if sensibilidad.empty:
        return

    resumen = (
        sensibilidad.groupby(["NivelDemanda", "CostePorKm"], as_index=False)
        .agg(
            CosteTotalMedio=("CosteTotalGA", "mean"),
            DistanciaMediaKm=("DistanciaKmGA", "mean"),
            ProductoresMedia=("NumeroProductoresGA", "mean"),
        )
    )

    fig, ax = plt.subplots(figsize=(8, 5))
    for nivel in ["Baja", "Media", "Alta"]:
        sub = resumen[resumen["NivelDemanda"] == nivel].sort_values("CostePorKm")
        if not sub.empty:
            ax.plot(
                sub["CostePorKm"],
                sub["CosteTotalMedio"],
                marker="o",
                label=nivel,
            )
    ax.set_title("Sensibilidad del coste total al coste por km")
    ax.set_xlabel("Coste de transporte (€/km)")
    ax.set_ylabel("Coste total medio (€)")
    ax.legend()
    ax.grid(True, alpha=0.25)
    guardar_figura(fig, "09_sensibilidad_coste_total.png")

    fig, ax = plt.subplots(figsize=(8, 5))
    for nivel in ["Baja", "Media", "Alta"]:
        sub = resumen[resumen["NivelDemanda"] == nivel].sort_values("CostePorKm")
        if not sub.empty:
            ax.plot(
                sub["CostePorKm"],
                sub["DistanciaMediaKm"],
                marker="o",
                label=nivel,
            )
    ax.set_title("Sensibilidad de la distancia al coste por km")
    ax.set_xlabel("Coste de transporte (€/km)")
    ax.set_ylabel("Distancia media (km)")
    ax.legend()
    ax.grid(True, alpha=0.25)
    guardar_figura(fig, "10_sensibilidad_distancia.png")


def generar_todas_las_graficas(
    semanas_rep: pd.DataFrame,
    mejores_por_semana: dict,
    ejecuciones: pd.DataFrame,
    resumen_semanas: pd.DataFrame,
    comparacion_baseline: pd.DataFrame,
    comparacion_milp: pd.DataFrame,
    sensibilidad: pd.DataFrame,
):
    if not GENERAR_GRAFICAS:
        return

    print("\nGenerando gráficas...")
    grafica_convergencia_representativas(semanas_rep, mejores_por_semana)
    grafica_estabilidad(ejecuciones)
    grafica_mejora_baseline(comparacion_baseline)
    grafica_costes_ga_baseline(comparacion_baseline)
    grafica_gap_milp(comparacion_milp)
    grafica_tiempos_ga_milp(comparacion_milp)
    grafica_productos_distancia(resumen_semanas)
    grafica_componentes_coste(ejecuciones)
    grafica_sensibilidad(sensibilidad)


def cargar_coordenadas_opcionales(ruta: Path):
    if not ruta.exists():
        return None

    try:
        prod = pd.read_excel(ruta, sheet_name="Productores")
        hubs = pd.read_excel(ruta, sheet_name="HUBs")
    except Exception:
        return None

    necesarias = {"Id", "Latitud", "Longitud"}
    if not necesarias.issubset(prod.columns) or not necesarias.issubset(hubs.columns):
        return None

    prod = prod[["Id", "Latitud", "Longitud"]].copy()
    hubs = hubs[["Id", "Latitud", "Longitud"]].copy()
    prod["Id"] = prod["Id"].map(normalizar_id)
    hubs["Id"] = hubs["Id"].map(normalizar_id)

    coords = {}
    for fila in pd.concat([prod, hubs], ignore_index=True).itertuples(index=False):
        try:
            coords[str(fila.Id)] = (float(fila.Latitud), float(fila.Longitud))
        except (TypeError, ValueError):
            continue

    return coords


def generar_mapas_representativos(
    semanas_rep: pd.DataFrame,
    mejores_por_semana: dict,
):
    if not GENERAR_MAPAS:
        return

    coords = cargar_coordenadas_opcionales(RUTA_DATOS_GEOGRAFICOS)
    if coords is None:
        print("\nMapas omitidos: no se encontraron coordenadas válidas.")
        return

    try:
        import folium
    except ImportError:
        print("\nMapas omitidos: instala folium para generarlos.")
        return

    print("\nGenerando mapas representativos con líneas rectas...")

    for _, fila in semanas_rep.iterrows():
        clave = (int(fila["Escenario"]), int(fila["NumeroSemanaSimulada"]))
        if clave not in mejores_por_semana:
            continue

        info = mejores_por_semana[clave]
        datos = info["datos"]
        mejor = info["resultado"].mejor
        activos = mejor.activos()
        ruta = [datos.hub] + activos + [datos.hub]

        if any(nodo not in coords for nodo in ruta):
            faltantes = [nodo for nodo in ruta if nodo not in coords]
            print(
                f"  Mapa omitido {fila['EtiquetaSemana']}: "
                f"faltan coordenadas para {faltantes}"
            )
            continue

        print(
            f"  {fila['EtiquetaSemana']} | {fila['NivelDemanda']} | "
            f"{len(activos)} productores"
        )

        latitudes = [coords[n][0] for n in ruta[:-1]]
        longitudes = [coords[n][1] for n in ruta[:-1]]

        mapa = folium.Map(
            location=[float(np.mean(latitudes)), float(np.mean(longitudes))],
            zoom_start=7,
            tiles="OpenStreetMap",
        )

        folium.Marker(
            location=coords[datos.hub],
            tooltip=f"HUB {datos.hub}",
            icon=folium.Icon(color="red", icon="home"),
        ).add_to(mapa)

        for orden_visita, productor in enumerate(activos, start=1):
            folium.Marker(
                location=coords[productor],
                tooltip=f"{orden_visita}. {productor}",
                icon=folium.Icon(color="blue", icon="info-sign"),
            ).add_to(mapa)


        geometria_recta = [coords[nodo] for nodo in ruta]
        folium.PolyLine(
            geometria_recta,
            weight=4,
            opacity=0.8,
            tooltip="Orden de visita (líneas rectas; no geometría vial)",
        ).add_to(mapa)

        nombre = (
            f"ruta_{sanitizar_nombre_archivo(str(fila['NivelDemanda']))}_"
            f"{fila['EtiquetaSemana']}.html"
        )
        ruta_salida_mapa = CARPETA_MAPAS / nombre
        mapa.save(str(ruta_salida_mapa))
        print(f"    Guardado: {ruta_salida_mapa}")


def dataframe_configuracion() -> pd.DataFrame:
    filas = [
        ("ModoEjecucion", MODO_EJECUCION),
        ("CostePorKmBase", COSTE_POR_KM_BASE),
        ("TratamientoProductosSinProveedor", "Excluir solo de optimizacion y conservar trazabilidad"),
        ("CriterioProveedorValido", "Producto con al menos un productor presente en precios y matriz de distancias"),
        ("TamanoPoblacion", TAMANO_POBLACION),
        ("NumeroGeneraciones", NUMERO_GENERACIONES),
        ("Paciencia", PACIENCIA),
        ("ProbabilidadCruce", PROBABILIDAD_CRUCE),
        ("ProbabilidadMutacion", PROBABILIDAD_MUTACION),
        ("PorcentajeElitismo", PORCENTAJE_ELITISMO),
        ("MaxExtrasIniciales", MAX_EXTRAS_INICIALES),
        ("SemillasGA", ", ".join(map(str, SEMILLAS_GA))),
        ("SemillaSeleccionSemanas", SEMILLA_SELECCION_SEMANAS),
        ("NSemanasBaja", N_SEMANAS_BAJA),
        ("NSemanasMedia", N_SEMANAS_MEDIA),
        ("NSemanasAlta", N_SEMANAS_ALTA),
        ("EjecutarMILP", EJECUTAR_MILP),
        ("NMILPPorNivel", N_MILP_POR_NIVEL),
        ("TimeLimitMILPSegundos", TIME_LIMIT_MILP_SEGUNDOS),
        ("MILPGapRelObjetivo", MILP_GAP_REL_OBJETIVO),
        ("MILPGapObjetivoPct", MILP_GAP_REL_OBJETIVO * 100.0),
        ("MILPTimeoutDuroExterno", MILP_USAR_TIMEOUT_DURO),
        ("MILPSolver", MILP_SOLVER),
        ("MILPMargenCierreSolverSegundos", MILP_MARGEN_CIERRE_SOLVER_SEGUNDOS),
        ("MILPRandomSeed", MILP_RANDOM_SEED),
        ("MILPHiGHSParallel", MILP_HIGHS_PARALLEL),
        ("MILPWatchdogIntervaloSegundos", MILP_WATCHDOG_INTERVALO_SEGUNDOS),
        ("FormulacionSubtoursMILP", FORMULACION_SUBTOURS_MILP),
        ("MILPThreads", MILP_THREADS),
        ("GenerarMapas", GENERAR_MAPAS),
        ("MapasModoTrazado", MAPAS_MODO_TRAZADO),
        ("EjecutarSensibilidad", EJECUTAR_SENSIBILIDAD),
        ("CostesKmSensibilidad", ", ".join(map(str, COSTES_KM_SENSIBILIDAD))),
        ("Python", sys.version.split()[0]),
        ("Sistema", platform.platform()),
    ]
    return pd.DataFrame(filas, columns=["Parametro", "Valor"])


def formatear_excel_salida(ruta: Path):
    try:
        from openpyxl import load_workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        return

    wb = load_workbook(ruta)

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(horizontal="center", vertical="center")


        for col_cells in ws.columns:
            max_len = 0
            col_idx = col_cells[0].column
            for cell in col_cells[:200]:
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 10), 35)

    wb.save(ruta)


def guardar_excel_final(
    configuracion: pd.DataFrame,
    coherencia: pd.DataFrame,
    productos_sin_proveedor: pd.DataFrame,
    resumen_todas_semanas: pd.DataFrame,
    umbrales: pd.DataFrame,
    semanas: pd.DataFrame,
    ejecuciones: pd.DataFrame,
    resumen_semanas: pd.DataFrame,
    resumen_niveles: pd.DataFrame,
    baseline: pd.DataFrame,
    comparacion_baseline: pd.DataFrame,
    milp: pd.DataFrame,
    milp_asignacion: pd.DataFrame,
    milp_arcos: pd.DataFrame,
    comparacion_milp: pd.DataFrame,
    tiempos_1pct: pd.DataFrame,
    historial: pd.DataFrame,
    sensibilidad: pd.DataFrame,
    rutas_mejores: pd.DataFrame,
    compras_mejores: pd.DataFrame,
    resumen_final: pd.DataFrame,
):
    print("\nGuardando Excel final...")

    with pd.ExcelWriter(RUTA_EXCEL_SALIDA, engine="openpyxl") as writer:
        configuracion.to_excel(writer, sheet_name="Configuracion", index=False)
        coherencia.to_excel(writer, sheet_name="CoherenciaDemandaOferta", index=False)
        productos_sin_proveedor.to_excel(writer, sheet_name="ProductosSinProveedor", index=False)
        resumen_todas_semanas.to_excel(writer, sheet_name="ResumenTodasSemanas", index=False)
        umbrales.to_excel(writer, sheet_name="UmbralesDemanda", index=False)
        semanas.to_excel(writer, sheet_name="SemanasSeleccionadas", index=False)
        ejecuciones.to_excel(writer, sheet_name="EjecucionesGA", index=False)
        resumen_semanas.to_excel(writer, sheet_name="ResumenSemanas", index=False)
        resumen_niveles.to_excel(writer, sheet_name="ResumenNiveles", index=False)
        baseline.to_excel(writer, sheet_name="Baseline", index=False)
        comparacion_baseline.to_excel(writer, sheet_name="ComparacionGABaseline", index=False)
        milp.to_excel(writer, sheet_name="MILP", index=False)
        milp_asignacion.to_excel(writer, sheet_name="MILP_Asignacion", index=False)
        milp_arcos.to_excel(writer, sheet_name="MILP_Arcos", index=False)
        comparacion_milp.to_excel(writer, sheet_name="ComparacionGAMILP", index=False)
        tiempos_1pct.to_excel(writer, sheet_name="TiempoHasta1PctMILP", index=False)
        historial.to_excel(writer, sheet_name="Convergencia", index=False)
        sensibilidad.to_excel(writer, sheet_name="SensibilidadCosteKm", index=False)
        rutas_mejores.to_excel(writer, sheet_name="RutasMejores", index=False)
        compras_mejores.to_excel(writer, sheet_name="ComprasMejores", index=False)
        resumen_final.to_excel(writer, sheet_name="ResumenFinal", index=False)

    formatear_excel_salida(RUTA_EXCEL_SALIDA)


def main():
    print("=" * 78)
    print("TFM - OPTIMIZACIÓN Y VALIDACIÓN COMPLETA DEL ALGORITMO GENÉTICO")
    print("=" * 78)

    asegurar_carpetas()
    maestros = cargar_datos_maestros()


    resumen_todas = crear_resumen_semanal(maestros.demanda_df)
    q33 = resumen_todas.attrs["Q33"]
    q67 = resumen_todas.attrs["Q67"]

    semanas = seleccionar_semanas_experimento(
        resumen_todas,
        SEMILLA_SELECCION_SEMANAS,
    )

    print("\nUmbrales de demanda semanal OPTIMIZABLE:")
    print(f"Q33 = {q33:.3f}")
    print(f"Q67 = {q67:.3f}")
    print("\nSemanas seleccionadas por nivel:")
    print(semanas["NivelDemanda"].value_counts())

    n_ejecuciones = len(semanas) * len(SEMILLAS_GA)
    print(f"\nEjecuciones GA previstas: {n_ejecuciones}")


    semanas_rep = seleccionar_representativas_por_nivel(
        semanas,
        N_SEMANAS_SENSIBILIDAD_POR_NIVEL,
    )


    (
        ejecuciones,
        historial,
        mejores_por_semana,
        rutas_mejores,
        compras_mejores,
    ) = ejecutar_bateria_ga(maestros, semanas)

    resumen_semanas = crear_resumen_ga_por_semana(ejecuciones)
    resumen_niveles = crear_resumen_por_nivel(resumen_semanas)


    baseline = ejecutar_baselines(maestros, semanas)
    comparacion_baseline = comparar_ga_baseline(resumen_semanas, baseline)


    semanas_milp = seleccionar_representativas_por_nivel(
        semanas,
        N_MILP_POR_NIVEL,
    )

    milp, milp_asignacion, milp_arcos = ejecutar_milps(
        maestros,
        semanas_milp,
    )

    comparacion_milp, tiempos_1pct = comparar_ga_milp(
        resumen_semanas,
        ejecuciones,
        historial,
        milp,
    )


    sensibilidad = ejecutar_sensibilidad(
        maestros,
        semanas_rep,
    )


    resumen_final = crear_resumen_final(
        ejecuciones,
        resumen_semanas,
        comparacion_baseline,
        comparacion_milp,
        tiempos_1pct,
        maestros.resumen_coherencia_df,
    )

    configuracion = dataframe_configuracion()
    umbrales = pd.DataFrame({
        "Indicador": [
            "Q33 demanda semanal optimizable",
            "Q67 demanda semanal optimizable",
            "Definición Baja",
            "Definición Media",
            "Definición Alta",
        ],
        "Valor": [
            q33,
            q67,
            f"DemandaTotalOptimizable <= {q33:.6f}",
            f"{q33:.6f} < DemandaTotalOptimizable <= {q67:.6f}",
            f"DemandaTotalOptimizable > {q67:.6f}",
        ],
    })


    generar_todas_las_graficas(
        semanas_rep,
        mejores_por_semana,
        ejecuciones,
        resumen_semanas,
        comparacion_baseline,
        comparacion_milp,
        sensibilidad,
    )

    generar_mapas_representativos(
        semanas_rep,
        mejores_por_semana,
    )


    guardar_excel_final(
        configuracion=configuracion,
        coherencia=maestros.resumen_coherencia_df,
        productos_sin_proveedor=maestros.productos_sin_proveedor_df,
        resumen_todas_semanas=resumen_todas,
        umbrales=umbrales,
        semanas=semanas,
        ejecuciones=ejecuciones,
        resumen_semanas=resumen_semanas,
        resumen_niveles=resumen_niveles,
        baseline=baseline,
        comparacion_baseline=comparacion_baseline,
        milp=milp,
        milp_asignacion=milp_asignacion,
        milp_arcos=milp_arcos,
        comparacion_milp=comparacion_milp,
        tiempos_1pct=tiempos_1pct,
        historial=historial,
        sensibilidad=sensibilidad,
        rutas_mejores=rutas_mejores,
        compras_mejores=compras_mejores,
        resumen_final=resumen_final,
    )

    print("\n" + "=" * 78)
    print("EXPERIMENTO FINALIZADO")
    print("=" * 78)
    print("\nResumen final:")
    print(resumen_final.to_string(index=False))
    print("\nExcel:")
    print(RUTA_EXCEL_SALIDA)
    print("\nGráficas:")
    print(CARPETA_GRAFICAS)
    print("\nMapas:")
    print(CARPETA_MAPAS)


if __name__ == "__main__":
    main()
